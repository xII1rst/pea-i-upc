"""PEA-i UPC: interfaz Tkinter conectada al backend C++.

Ejecutar: python src/python/Taller2_REMR.py [--data-dir CARPETA]
Con --python-backend se usa la implementación Python independiente de estructuras enlazadas.
La lógica de datos puede probarse sin un servidor gráfico.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
import hashlib
from html.parser import HTMLParser
import http.client
import io
import ipaddress
import json
import os
from pathlib import Path
import queue
import re
import shutil
import socket
import subprocess
import tempfile
import threading
from typing import Any, Callable, Iterator, NamedTuple
import unicodedata
from urllib.parse import parse_qs, urlencode, urlparse
from urllib.request import (HTTPRedirectHandler, HTTPHandler, HTTPSHandler,
                            ProxyHandler, Request, build_opener)
import uuid
import zipfile


SCHEMA_VERSION = "2"
ENTITY_FIELDS = {
    "grupos": ("id", "nombre", "codigo_gruplac", "fecha_creacion", "unidad", "responsable", "categoria", "descripcion", "objetivos", "mision", "vision", "lineas", "url", "fuente", "activo"),
    "investigadores": ("id", "nombre", "codigo_cvlac", "afiliacion", "categoria", "contacto", "url", "fuente", "activo"),
    "productos": ("id", "titulo", "anio", "fecha", "familia", "tipologia", "categoria", "validacion", "observacion", "doi", "url", "fuente", "activo"),
    "planes": ("id", "grupo_id", "nombre", "inicio", "fin", "objetivo", "indicador", "meta", "actividad", "activo"),
}
RELATION_FIELDS = {
    "membresias": ("grupo_id", "investigador_id", "rol", "inicio", "fin", "activo"),
    "autorias": ("producto_id", "investigador_id", "orden", "rol", "activo"),
    "grupos_productos": ("grupo_id", "producto_id", "origen", "activo"),
}
RELATION_ENDS = {
    "membresias": ("grupo_id", "investigador_id", "grupos", "investigadores"),
    "autorias": ("producto_id", "investigador_id", "productos", "investigadores"),
    "grupos_productos": ("grupo_id", "producto_id", "grupos", "productos"),
}
ALL_FIELDS = {**ENTITY_FIELDS, **RELATION_FIELDS}
PAGE_SIZE = 100
REQUIRED = {
    "grupos": ("id", "nombre"), "investigadores": ("id", "nombre"),
    "productos": ("id", "titulo"), "planes": ("id", "grupo_id", "nombre"),
}
LABELS = {
    "id": "ID", "nombre": "Nombre", "codigo_gruplac": "Código GrupLAC",
    "fecha_creacion": "Fecha de creación", "unidad": "Unidad académica", "responsable": "Responsable",
    "categoria": "Categoría", "descripcion": "Descripción", "objetivos": "Objetivos",
    "mision": "Misión", "vision": "Visión", "lineas": "Líneas", "url": "URL",
    "fuente": "Fuente", "activo": "Activo", "codigo_cvlac": "Código CvLAC",
    "afiliacion": "Afiliación", "contacto": "Contacto público", "titulo": "Título",
    "anio": "Año", "fecha": "Fecha", "familia": "Familia", "tipologia": "Tipología",
    "validacion": "Validación", "observacion": "Observación", "doi": "DOI",
    "grupo_id": "Grupo (ID)", "investigador_id": "Investigador (ID)",
    "producto_id": "Producto (ID)", "inicio": "Inicio", "fin": "Fin",
    "objetivo": "Objetivo", "indicador": "Indicador", "meta": "Meta",
    "actividad": "Actividad", "rol": "Rol", "orden": "Orden",
    "origen": "Origen de asociación",
}
LONG_FIELDS = {"descripcion", "objetivos", "mision", "vision", "lineas", "observacion", "objetivo", "indicador", "meta", "actividad", "fuente"}
MAX_FIELD_LEN = 100_000
MAX_CSV_ROWS = 500_000
MAX_PDF_TEXT = 5_000_000
MAX_PDF_MEMORY = 512 * 1024 * 1024
PDF_TIMEOUT = 25
MAX_LOCAL_FILE = 32 * 1024 * 1024
MAX_ARCHIVE_UNPACKED = 128 * 1024 * 1024


class DataError(ValueError):
    """Dato inválido o relación que rompería la integridad."""


class ListNode:
    def __init__(self, value: dict[str, str]):
        self.value = value
        self.prev: ListNode | None = None
        self.next: ListNode | None = None


class DoublyLinkedList:
    """Lista fuente de verdad para entidades; índices auxiliares no la reemplazan."""

    def __init__(self) -> None:
        self.head: ListNode | None = None
        self.tail: ListNode | None = None
        self.length = 0
        self.by_id: dict[str, ListNode] = {}

    def append(self, value: dict[str, str]) -> None:
        key = value["id"]
        if key in self.by_id:
            raise DataError(f"ID duplicado: {key}")
        node = ListNode(value)
        node.prev = self.tail
        if self.tail:
            self.tail.next = node
        else:
            self.head = node
        self.tail = node
        self.by_id[key] = node
        self.length += 1

    def get(self, key: str) -> dict[str, str] | None:
        node = self.by_id.get(key)
        return node.value if node else None

    def remove(self, key: str) -> dict[str, str]:
        node = self.by_id.pop(key, None)
        if node is None:
            raise DataError(f"No existe ID {key}")
        if node.prev:
            node.prev.next = node.next
        else:
            self.head = node.next
        if node.next:
            node.next.prev = node.prev
        else:
            self.tail = node.prev
        self.length -= 1
        return node.value

    def __iter__(self) -> Iterator[dict[str, str]]:
        node = self.head
        while node:
            yield node.value
            node = node.next

    def backwards(self) -> Iterator[dict[str, str]]:
        node = self.tail
        while node:
            yield node.value
            node = node.prev


class RelationNode:
    def __init__(self, left: str, right: str, value: dict[str, str]):
        self.left = left
        self.right = right
        self.value = value
        self.next_left: RelationNode | None = None
        self.next_right: RelationNode | None = None


class MultiList:
    """Cada vínculo pertenece simultáneamente a dos cadenas de recorrido."""

    def __init__(self, left_field: str, right_field: str) -> None:
        self.left_field = left_field
        self.right_field = right_field
        self.left_heads: dict[str, RelationNode] = {}
        self.right_heads: dict[str, RelationNode] = {}
        self.by_pair: dict[tuple[str, str], RelationNode] = {}

    def add(self, row: dict[str, str]) -> None:
        left, right = row[self.left_field], row[self.right_field]
        pair = (left, right)
        if pair in self.by_pair:
            raise DataError(f"Relación duplicada: {left} / {right}")
        node = RelationNode(left, right, row)
        node.next_left = self.left_heads.get(left)
        node.next_right = self.right_heads.get(right)
        self.left_heads[left] = node
        self.right_heads[right] = node
        self.by_pair[pair] = node

    def get(self, left: str, right: str) -> dict[str, str] | None:
        node = self.by_pair.get((left, right))
        return node.value if node else None

    def by_left(self, left: str) -> Iterator[dict[str, str]]:
        node = self.left_heads.get(left)
        while node:
            yield node.value
            node = node.next_left

    def by_right(self, right: str) -> Iterator[dict[str, str]]:
        node = self.right_heads.get(right)
        while node:
            yield node.value
            node = node.next_right

    def remove(self, left: str, right: str) -> dict[str, str]:
        node = self.by_pair.pop((left, right), None)
        if not node:
            raise DataError("No existe esa relación")
        previous = None
        current = self.left_heads[left]
        while current is not node:
            previous, current = current, current.next_left
        if previous:
            previous.next_left = node.next_left
        elif node.next_left:
            self.left_heads[left] = node.next_left
        else:
            del self.left_heads[left]
        previous = None
        current = self.right_heads[right]
        while current is not node:
            previous, current = current, current.next_right
        if previous:
            previous.next_right = node.next_right
        elif node.next_right:
            self.right_heads[right] = node.next_right
        else:
            del self.right_heads[right]
        return node.value

    def __iter__(self) -> Iterator[dict[str, str]]:
        for left in self.left_heads:
            yield from self.by_left(left)


class StackNode:
    def __init__(self, value: str, next_node: StackNode | None = None):
        self.value, self.next = value, next_node


class LinkedStack:
    def __init__(self, limit: int = 30) -> None:
        self.top: StackNode | None = None
        self.length = 0
        self.limit = limit

    def push(self, value: str) -> None:
        self.top = StackNode(value, self.top)
        self.length += 1
        if self.length > self.limit:
            node = self.top
            for _ in range(self.limit - 1):
                node = node.next  # type: ignore[assignment]
            node.next = None
            self.length = self.limit

    def peek(self) -> str | None:
        return self.top.value if self.top else None

    def pop(self) -> str | None:
        if not self.top:
            return None
        value = self.top.value
        self.top = self.top.next
        self.length -= 1
        return value

    def clear(self) -> None:
        self.top, self.length = None, 0

    def __iter__(self) -> Iterator[str]:
        node = self.top
        while node:
            yield node.value
            node = node.next


class QueueNode:
    def __init__(self, value: dict[str, str]):
        self.value = value
        self.next: QueueNode | None = None


class LinkedQueue:
    def __init__(self) -> None:
        self.front: QueueNode | None = None
        self.rear: QueueNode | None = None
        self.length = 0

    def enqueue(self, value: dict[str, str]) -> None:
        node = QueueNode(value)
        if self.rear:
            self.rear.next = node
        else:
            self.front = node
        self.rear = node
        self.length += 1

    def prepend(self, value: dict[str, str]) -> None:
        node = QueueNode(value)
        node.next = self.front
        self.front = node
        if self.rear is None:
            self.rear = node
        self.length += 1

    def remove(self, job_id: str) -> None:
        previous = None
        node = self.front
        while node and node.value["id"] != job_id:
            previous, node = node, node.next
        if node is None:
            raise DataError("Trabajo de cola no encontrado")
        if previous:
            previous.next = node.next
        else:
            self.front = node.next
        if self.rear is node:
            self.rear = previous
        self.length -= 1

    def peek(self) -> dict[str, str] | None:
        return self.front.value if self.front else None

    def dequeue(self) -> dict[str, str] | None:
        if not self.front:
            return None
        value = self.front.value
        self.front = self.front.next
        if not self.front:
            self.rear = None
        self.length -= 1
        return value

    def clear(self) -> None:
        self.front, self.rear, self.length = None, None, 0

    def __iter__(self) -> Iterator[dict[str, str]]:
        node = self.front
        while node:
            yield node.value
            node = node.next


def clean_row(kind: str, values: dict[str, Any]) -> dict[str, str]:
    if kind not in ALL_FIELDS:
        raise DataError(f"Tipo desconocido: {kind}")
    extra = set(values) - set(ALL_FIELDS[kind])
    if extra:
        raise DataError(f"Columnas desconocidas: {', '.join(sorted(extra))}")
    row = {field: str(values.get(field, "") or "").strip() for field in ALL_FIELDS[kind]}
    row["activo"] = row["activo"] or "1"
    if row["activo"] not in ("0", "1"):
        raise DataError("Activo debe ser 0 o 1")
    for field in REQUIRED.get(kind, ()):
        if not row[field]:
            raise DataError(f"Falta {LABELS[field]}")
    for field in ALL_FIELDS[kind]:
        if field.endswith("_id") and not row[field]:
            raise DataError(f"Falta {LABELS[field]}")
    if kind == "productos":
        if row["fecha"] and not row["anio"]:
            row["anio"] = row["fecha"][:4]
        if row["anio"]:
            try:
                year = int(row["anio"])
            except ValueError as exc:
                raise DataError("Año inválido") from exc
            if not 1900 <= year <= date.today().year + 1:
                raise DataError("Año fuera de rango")
    if kind == "productos":
        row["validacion"] = row["validacion"] or "pendiente"
        if row["validacion"] not in ("pendiente", "validado", "rechazado"):
            raise DataError("Validación: pendiente, validado o rechazado")
    if kind == "autorias" and row["orden"]:
        if not row["orden"].isdigit() or int(row["orden"]) < 1:
            raise DataError("Orden de autor debe ser positivo")
    for field in ("fecha", "fecha_creacion", "inicio", "fin"):
        if field in row and row[field]:
            try:
                date.fromisoformat(row[field])
            except ValueError as exc:
                raise DataError(f"{LABELS[field]} debe ser AAAA-MM-DD") from exc
    if kind == "productos" and row["fecha"] and row["anio"]:
        if row["fecha"][:4] != row["anio"]:
            raise DataError("El año y la fecha del producto no coinciden")
    if kind in ("planes", "membresias") and row["inicio"] and row["fin"]:
        if row["inicio"] > row["fin"]:
            raise DataError("La fecha final precede la inicial")
    for field, value in row.items():
        if len(value) > MAX_FIELD_LEN:
            raise DataError(f"El campo {LABELS.get(field, field)} supera el tamaño permitido")
    return row


def _legacy_row(kind: str, values: dict[str, str]) -> dict[str, str]:
    """Adapta filas de la versión 1 sin alterar los datos persistidos de origen."""
    row = values.copy()
    if kind == "grupos":
        row.pop("sigla", None)
    if kind == "productos":
        row.setdefault("validacion", "pendiente")
    return row


class Repository:
    """Operaciones del dominio. Las entidades viven en listas; relaciones en multilistas."""

    def __init__(self) -> None:
        self.entities = {name: DoublyLinkedList() for name in ENTITY_FIELDS}
        self.relations = {
            name: MultiList(ends[0], ends[1]) for name, ends in RELATION_ENDS.items()
        }
        self.history = LinkedStack()
        self.queue = LinkedQueue()
        self.dirty = False

    def rows(self, kind: str) -> list[dict[str, str]]:
        source = self.entities.get(kind) or self.relations.get(kind)
        if source is None:
            raise DataError(f"Tipo desconocido: {kind}")
        rows = [row.copy() for row in source]
        if kind in RELATION_ENDS:
            left, right = RELATION_ENDS[kind][:2]
            rows.sort(key=lambda row: (row[left], row[right]))
        return rows

    def page(self, kind: str, offset: int = 0, limit: int = PAGE_SIZE,
             query: str = "") -> dict[str, Any]:
        source = self.entities.get(kind) or self.relations.get(kind)
        if source is None:
            raise DataError(f"Tipo desconocido: {kind}")
        needle = query.casefold().strip()
        rows = (row for row in source if not needle or needle in " ".join(row.values()).casefold())
        if kind in RELATION_ENDS:
            left, right = RELATION_ENDS[kind][:2]
            matches = sorted(rows, key=lambda row: (row[left], row[right]))
            return {"rows": [row.copy() for row in matches[offset:offset + limit]], "total": len(matches)}
        page_rows = []
        total = 0
        for row in rows:
            if offset <= total < offset + limit:
                page_rows.append(row.copy())
            total += 1
        return {"rows": page_rows, "total": total}

    def summary(self) -> dict[str, Any]:
        categories: dict[str, int] = {}
        for product in self.entities["productos"]:
            if product["activo"] == "1":
                name = product["categoria"] or "Sin dato"
                categories[name] = categories.get(name, 0) + 1
        return {"active_groups": sum(row["activo"] == "1" for row in self.entities["grupos"]),
                "active_people": sum(row["activo"] == "1" for row in self.entities["investigadores"]),
                "categories": categories}

    def get(self, kind: str, key: str | tuple[str, str]) -> dict[str, str] | None:
        if kind in ENTITY_FIELDS:
            return self.entities[kind].get(str(key))
        if kind in RELATION_FIELDS and isinstance(key, tuple):
            return self.relations[kind].get(*key)
        return None

    def related(self, kind: str, key: str, side: str) -> list[dict[str, str]]:
        relation = self.relations[kind]
        return list(relation.by_left(key) if side == "left" else relation.by_right(key))

    def related_page(self, kind: str, key: str, side: str, offset: int = 0,
                     limit: int = PAGE_SIZE, active_only: bool = False) -> dict[str, Any]:
        relation = self.relations[kind]
        source = relation.by_left(key) if side == "left" else relation.by_right(key)
        rows = []
        total = 0
        for row in source:
            if active_only and row["activo"] != "1":
                continue
            if offset <= total < offset + limit:
                rows.append(row.copy())
            total += 1
        return {"rows": rows, "total": total}

    def history_size(self) -> int:
        return self.history.length

    def suggest_id(self, kind: str) -> str:
        return new_id({"grupos": "G", "investigadores": "I",
                       "productos": "P", "planes": "PL"}[kind])

    def snapshot(self) -> str:
        state = {kind: self.rows(kind) for kind in ALL_FIELDS}
        state["cola_validacion"] = [job.copy() for job in self.queue]
        return json.dumps(state, ensure_ascii=False, separators=(",", ":"))

    def _restore(self, snapshot: str) -> None:
        state = json.loads(snapshot)
        fresh = Repository()
        for kind in ENTITY_FIELDS:
            for row in state[kind]:
                fresh.create(kind, _legacy_row(kind, row), remember=False)
        for kind in RELATION_FIELDS:
            for row in state[kind]:
                fresh.create(kind, row, remember=False)
        for job in state.get("cola_validacion", []):
            fresh._validate_job(job)
            fresh.queue.enqueue(job)
        self.entities, self.relations, self.queue = fresh.entities, fresh.relations, fresh.queue
        self.dirty = True

    def _remember(self) -> None:
        self.history.push(self.snapshot())
        self.dirty = True

    @staticmethod
    def _decode_delta(snapshot: str) -> list[dict[str, str]] | None:
        state = json.loads(snapshot)
        if "__undo" not in state:
            return None
        if set(state) != {"__undo"} or not isinstance(state["__undo"], list) or not state["__undo"]:
            raise DataError("Historial de cambios inválido")
        for change in state["__undo"]:
            if not isinstance(change, dict):
                raise DataError("Historial de cambios inválido")
            op, kind = change.get("__op"), change.get("__kind")
            if op == "queue_remove":
                if not change.get("id"):
                    raise DataError("Historial de cola inválido")
            elif op == "queue_prepend":
                if not {"id", "producto_id", "motivo", "creado"} <= set(change):
                    raise DataError("Historial de cola inválido")
            elif op in ("create", "replace") and kind in ALL_FIELDS:
                clean_row(kind, _legacy_row(kind, {key: value for key, value in change.items()
                                                   if not key.startswith("__")}))
            elif op == "delete" and kind in ALL_FIELDS:
                keys = ("id",) if kind in ENTITY_FIELDS else RELATION_ENDS[kind][:2]
                if any(not change.get(key) for key in keys):
                    raise DataError("Historial sin clave")
            else:
                raise DataError("Operación de historial desconocida")
        return state["__undo"]

    def _apply_delta(self, changes: list[dict[str, str]]) -> None:
        for change in changes:
            op, kind = change["__op"], change["__kind"]
            if op == "queue_remove":
                self.queue.remove(change["id"])
            elif op == "queue_prepend":
                job = {name: change[name] for name in ("id", "producto_id", "motivo", "creado")}
                self._validate_job(job)
                self.queue.prepend(job)
            else:
                row = _legacy_row(kind, {key: value for key, value in change.items()
                                         if not key.startswith("__")})
                key = row["id"] if kind in ENTITY_FIELDS else tuple(row[name] for name in RELATION_ENDS[kind][:2])
                if op == "delete":
                    self.delete(kind, key, remember=False)
                elif op == "create":
                    self.create(kind, row, remember=False)
                else:
                    current = self.get(kind, key)
                    if current is None:
                        raise DataError("Registro de historial no encontrado")
                    current.update(row)
        self.dirty = True

    def undo(self) -> bool:
        previous = self.history.peek()
        if previous is None:
            return False
        changes = self._decode_delta(previous)
        if changes is None:
            self._restore(previous)
        else:
            self._apply_delta(changes)
        self.history.pop()
        return True

    def clear_history(self) -> None:
        self.history.clear()
        self.dirty = True

    def _check_links(self, kind: str, row: dict[str, str]) -> None:
        if kind == "planes" and not self.get("grupos", row["grupo_id"]):
            raise DataError("El grupo del plan no existe")
        if kind in RELATION_ENDS:
            left, right, left_kind, right_kind = RELATION_ENDS[kind]
            if not self.get(left_kind, row[left]) or not self.get(right_kind, row[right]):
                raise DataError("La relación apunta a una entidad inexistente")

    def _unique_external(self, kind: str, row: dict[str, str], old_id: str = "") -> None:
        field = {"grupos": "codigo_gruplac", "investigadores": "codigo_cvlac", "productos": "doi"}.get(kind)
        if not field or not row[field]:
            return
        normalized = row[field].strip().lower().removeprefix("https://doi.org/") if field == "doi" else row[field].strip().lower()
        for existing in self.entities[kind]:
            if existing["id"] == old_id or not existing[field]:
                continue
            candidate = existing[field].strip().lower().removeprefix("https://doi.org/") if field == "doi" else existing[field].strip().lower()
            if candidate == normalized:
                raise DataError(f"{LABELS[field]} ya registrado en {existing['id']}")

    def create(self, kind: str, values: dict[str, Any], remember: bool = True) -> dict[str, str]:
        row = clean_row(kind, values)
        self._check_links(kind, row)
        if kind in ENTITY_FIELDS:
            self._unique_external(kind, row)
            if self.get(kind, row["id"]):
                raise DataError("ID duplicado")
        else:
            left, right = RELATION_ENDS[kind][:2]
            if self.get(kind, (row[left], row[right])):
                raise DataError("Relación duplicada")
        if remember:
            self._remember()
        if kind in ENTITY_FIELDS:
            self.entities[kind].append(row)
        else:
            self.relations[kind].add(row)
        return row

    def update(self, kind: str, key: str | tuple[str, str], values: dict[str, Any], remember: bool = True) -> dict[str, str]:
        current = self.get(kind, key)
        if current is None:
            raise DataError("Registro no encontrado")
        updated = clean_row(kind, {**current, **values})
        fixed = ("id",) if kind in ENTITY_FIELDS else RELATION_ENDS[kind][:2]
        if any(updated[field] != current[field] for field in fixed):
            raise DataError("Los identificadores son estables; cree otra relación si cambian")
        self._check_links(kind, updated)
        if kind in ENTITY_FIELDS:
            self._unique_external(kind, updated, current["id"])
        if kind == "productos" and (current["validacion"] != updated["validacion"] or current["categoria"] != updated["categoria"]):
            if not updated["observacion"] or updated["observacion"] == current["observacion"]:
                raise DataError("Explique el cambio de categoría o validación en Observación")
        if remember:
            self._remember()
        current.update(updated)
        return current

    def delete(self, kind: str, key: str | tuple[str, str], remember: bool = True) -> None:
        row = self.get(kind, key)
        if not row:
            raise DataError("Registro no encontrado")
        if kind in ENTITY_FIELDS:
            if kind == "grupos" and any(p["grupo_id"] == key for p in self.entities["planes"]):
                raise DataError("El grupo tiene planes; elimínelos o desactívelo")
            for rel_kind, (left, right, left_kind, right_kind) in RELATION_ENDS.items():
                field = left if kind == left_kind else right if kind == right_kind else None
                if field and any(rel[field] == key for rel in self.relations[rel_kind]):
                    raise DataError("Hay relaciones asociadas; elimínelas o desactive el registro")
            if kind == "productos" and any(job["producto_id"] == key for job in self.queue):
                raise DataError("Hay revisiones pendientes para el producto")
        if remember:
            self._remember()
        if kind in ENTITY_FIELDS:
            self.entities[kind].remove(str(key))
        else:
            self.relations[kind].remove(*key)

    def toggle(self, kind: str, key: str | tuple[str, str]) -> None:
        row = self.get(kind, key)
        if not row:
            raise DataError("Registro no encontrado")
        self.update(kind, key, {"activo": "0" if row["activo"] == "1" else "1"})

    def product_scope(self, view: str = "Todos", selected_id: str = "") -> list[dict[str, str]]:
        if view == "Grupo":
            ids = {r["producto_id"] for r in self.relations["grupos_productos"].by_left(selected_id) if r["activo"] == "1"} if selected_id else set()
        elif view == "Investigador":
            ids = {r["producto_id"] for r in self.relations["autorias"].by_right(selected_id) if r["activo"] == "1"} if selected_id else set()
        elif view == "Producto":
            ids = {selected_id} if selected_id else set()
        else:
            ids = None
        return [p for p in self.entities["productos"] if p["activo"] == "1" and (ids is None or p["id"] in ids)]

    def statistics(self, view: str = "Todos", selected_id: str = "", start: int | None = None,
                   end: int | None = None, category: str = "", status: str = "",
                   offset: int = 0, limit: int | None = None) -> dict[str, Any]:
        # Hipercubo lógico: cada producto es un hecho; grupo e investigador llegan
        # por multilistas, y año, tipología, categoría y validación son dimensiones.
        # La vista elige grupo/investigador/producto; los filtros recortan el conjunto
        # y estos conteos se calculan bajo demanda. No hay cubo OLAP materializado
        # ni filtro simultáneo por grupo e investigador.
        if view not in ("Todos", "Grupo", "Investigador", "Producto"):
            raise DataError("Vista desconocida")
        if start is not None and end is not None and start > end:
            raise DataError("El año inicial supera al final")
        products = []
        total = 0
        tallies: dict[str, dict[str, int]] = {name: {} for name in ("anio", "tipologia", "categoria", "validacion")}
        for row in self.product_scope(view, selected_id):
            year = int(row["anio"]) if row["anio"] else None
            if start is not None and (year is None or year < start):
                continue
            if end is not None and (year is None or year > end):
                continue
            if category and row["categoria"] != category:
                continue
            if status and row["validacion"] != status:
                continue
            if total >= offset and (limit is None or len(products) < limit):
                products.append(row.copy())
            total += 1
            for name, counts in tallies.items():
                key = row[name] or "Sin dato"
                counts[key] = counts.get(key, 0) + 1
        return {"productos": products, "total": total,
                "por_anio": tallies["anio"], "por_tipologia": tallies["tipologia"],
                "por_categoria": tallies["categoria"], "por_validacion": tallies["validacion"]}


    def queue_rows(self) -> list[dict[str, str]]:
        return list(self.queue)
    def queue_page(self, offset: int = 0, limit: int = PAGE_SIZE) -> dict[str, Any]:
        rows = []
        for index, job in enumerate(self.queue):
            if offset <= index < offset + limit:
                rows.append(job.copy())
            if index >= offset + limit:
                break
        return {"rows": rows, "total": self.queue.length}
    def queue_front(self) -> dict[str, str] | None:
        return self.queue.peek()
    def queue_size(self) -> int:
        return self.queue.length
    def _validate_job(self, job: dict[str, str]) -> None:
        if set(job) != {"id", "producto_id", "motivo", "creado"}:
            raise DataError("Formato inválido de la cola")
        if not self.get("productos", job["producto_id"]):
            raise DataError("La cola referencia un producto inexistente")
        if not job["id"] or not job["creado"]:
            raise DataError("Trabajo de cola incompleto")
    def enqueue_review(self, product_id: str, reason: str) -> None:
        if not self.get("productos", product_id):
            raise DataError("Producto inexistente")
        if any(job["producto_id"] == product_id for job in self.queue):
            raise DataError("El producto ya está en la cola")
        job = {"id": uuid.uuid4().hex[:12], "producto_id": product_id,
               "motivo": reason.strip(), "creado": datetime.now().isoformat(timespec="seconds")}
        self._remember()
        self.queue.enqueue(job)
    def process_review(self, status: str, observation: str) -> dict[str, str]:
        job = self.queue.peek()
        if not job:
            raise DataError("La cola está vacía")
        if status not in ("validado", "rechazado", "pendiente"):
            raise DataError("Estado inválido")
        if not observation.strip():
            raise DataError("Debe registrar una observación")
        current = self.get("productos", job["producto_id"])
        if current is None:
            raise DataError("El producto ya no existe")
        clean_row("productos", {**current, "validacion": status, "observacion": observation})
        if current["validacion"] != status and current["observacion"] == observation.strip():
            raise DataError("Escriba una observación nueva para cambiar la validación")
        self._remember()
        self.update("productos", job["producto_id"], {"validacion": status, "observacion": observation}, remember=False)
        self.queue.dequeue()
        return job
    def discard_review(self) -> dict[str, str]:
        if not self.queue.peek():
            raise DataError("La cola está vacía")
        self._remember()
        return self.queue.dequeue()  # type: ignore[return-value]

def _storage_cell(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r", "'")) else value


def _csv_read(path: Path, fields: tuple[str, ...], *, legacy: bool = False,
              encoded: bool = False) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        actual = set(reader.fieldnames or ())
        expected = set(fields)
        old_groups = legacy and path.stem == "grupos" and actual == expected | {"sigla"}
        old_products = legacy and path.stem == "productos" and actual == expected - {"validacion"}
        if reader.fieldnames is None or not (actual == expected or old_groups or old_products):
            raise DataError(f"Encabezados incorrectos en {path.name}; se esperan: {', '.join(fields)}")
        rows = []
        for number, row in enumerate(reader, start=2):
            if None in row or any(value is None for value in row.values()):
                raise DataError(f"Fila {number} malformada en {path.name}")
            if len(rows) >= MAX_CSV_ROWS:
                raise DataError(f"Demasiadas filas en {path.name}")
            if encoded:
                row = {key: value[1:] if value.startswith("'") else value for key, value in row.items()}
            rows.append(_legacy_row(path.stem, row) if legacy else row)
        return rows


def _csv_write(path: Path, fields: tuple[str, ...], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: _storage_cell(row.get(field, "")) for field in fields})
        stream.flush()
        os.fsync(stream.fileno())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _check_local_file(path: Path, *, archive: bool = False) -> None:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise DataError(f"No se puede abrir el archivo: {exc}") from exc
    if size > MAX_LOCAL_FILE:
        raise DataError("El archivo supera 32 MB")
    if archive:
        try:
            with zipfile.ZipFile(path) as package:
                members = package.infolist()
                if len(members) > 10_000 or sum(item.file_size for item in members) > MAX_ARCHIVE_UNPACKED:
                    raise DataError("El documento comprimido supera los límites de extracción")
        except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError) as exc:
            raise DataError("El archivo no es un documento Office válido") from exc


def _xlsx_to_csv(source: Path, target: Path) -> Path:
    """Convierte la primera hoja de un .xlsx a CSV respetando los límites de importación."""
    _check_local_file(source, archive=True)
    try:
        import openpyxl
    except ImportError as exc:
        raise DataError("Para importar Excel instale openpyxl (pip install openpyxl)") from exc
    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise DataError(f"No se pudo leer el Excel: {exc}") from exc
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        header_cells = [_cell_text(cell) for cell in (header or ())]
        if not any(header_cells):
            raise DataError("El Excel no tiene encabezados")
        if any(not name for name in header_cells):
            raise DataError("El Excel tiene columnas sin encabezado")
        if any(len(name) > MAX_FIELD_LEN for name in header_cells):
            raise DataError("El Excel tiene un encabezado demasiado largo")
        with target.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream, lineterminator="\n")
            writer.writerow(header_cells)
            count = 1
            for record in rows:
                count += 1
                if count > MAX_CSV_ROWS:
                    raise DataError("El Excel tiene demasiadas filas")
                cells = [_cell_text(cell) for cell in record]
                if any(len(cell) > MAX_FIELD_LEN for cell in cells):
                    raise DataError("El Excel tiene un valor demasiado largo")
                writer.writerow(cells)
    except DataError:
        raise
    except Exception as exc:
        raise DataError(f"No se pudo leer el Excel: {exc}") from exc
    finally:
        workbook.close()
    return target


def _neutralize_cell(value: str) -> str:
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


def _csv_write_safe(path: Path, fields: tuple[str, ...], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: _neutralize_cell(row.get(field, "")) for field in fields})
        stream.flush()
        os.fsync(stream.fileno())


def export_repository_spreadsheet(repo: Repository | CppRepository, directory: Path) -> None:
    """Exporta los datos neutralizados para abrirlos en hojas de cálculo sin fórmulas."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    _csv_write_safe(directory / "manifest.csv", ("version", "guardado"),
                    [{"version": SCHEMA_VERSION, "guardado": datetime.now().isoformat(timespec="seconds")}])
    for kind, fields in ALL_FIELDS.items():
        _csv_write_safe(directory / f"{kind}.csv", fields, repo.rows(kind))


def save_repository(repo: Repository, directory: Path) -> None:
    """Prepara todos los CSV antes de sustituir archivos; conserva copia anterior."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".pea-save-", dir=directory) as temp_name:
        staged = Path(temp_name)
        _csv_write(staged / "manifest.csv", ("version", "guardado"),
                   [{"version": SCHEMA_VERSION, "guardado": datetime.now().isoformat(timespec="seconds")}])
        for kind, fields in ALL_FIELDS.items():
            _csv_write(staged / f"{kind}.csv", fields, repo.rows(kind))
        _csv_write(staged / "cola_validacion.csv", ("id", "producto_id", "motivo", "creado"), list(repo.queue))
        _csv_write(staged / "historial.csv", ("orden", "snapshot"),
                   [{"orden": str(index), "snapshot": state} for index, state in enumerate(repo.history)])
        files = ["manifest.csv", *(f"{kind}.csv" for kind in ALL_FIELDS), "cola_validacion.csv", "historial.csv"]
        backup = directory / ".backup"
        backup.mkdir(exist_ok=True)
        for name in files:
            if (directory / name).exists():
                shutil.copy2(directory / name, backup / name)
        for name in files:
            os.replace(staged / name, directory / name)
    repo.dirty = False


def load_repository(directory: Path) -> Repository:
    directory = Path(directory)
    manifest = _csv_read(directory / "manifest.csv", ("version", "guardado"), encoded=True)
    if len(manifest) != 1 or manifest[0]["version"] not in ("1", SCHEMA_VERSION):
        raise DataError("Versión de datos incompatible")
    legacy = manifest[0]["version"] == "1"
    repo = Repository()
    for kind in ENTITY_FIELDS:
        for number, row in enumerate(_csv_read(directory / f"{kind}.csv", ENTITY_FIELDS[kind],
                                               legacy=legacy, encoded=not legacy), start=2):
            try:
                repo.create(kind, row, remember=False)
            except DataError as exc:
                raise DataError(f"{kind}.csv, fila {number}: {exc}") from exc
    for kind in RELATION_FIELDS:
        for number, row in enumerate(_csv_read(directory / f"{kind}.csv", RELATION_FIELDS[kind],
                                               encoded=not legacy), start=2):
            try:
                repo.create(kind, row, remember=False)
            except DataError as exc:
                raise DataError(f"{kind}.csv, fila {number}: {exc}") from exc
    job_ids: set[str] = set()
    queue_path = directory / "cola_validacion.csv"
    if not queue_path.exists() and not legacy:
        raise DataError("Falta cola_validacion.csv")
    jobs = (_csv_read(queue_path, ("id", "producto_id", "motivo", "creado"), encoded=not legacy)
            if queue_path.exists() else [])
    for job in jobs:
        repo._validate_job(job)
        if job["id"] in job_ids:
            raise DataError(f"Trabajo de cola duplicado: {job['id']}")
        job_ids.add(job["id"])
        repo.queue.enqueue(job)
    history = _csv_read(directory / "historial.csv", ("orden", "snapshot"), encoded=not legacy)
    for entry in reversed(history):
        # Cada snapshot se verifica antes de incorporarlo al historial.
        validator = Repository()
        try:
            if validator._decode_delta(entry["snapshot"]) is None:
                validator._restore(entry["snapshot"])
        except (KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
            raise DataError("Historial dañado") from exc
        repo.history.push(entry["snapshot"])
    repo.dirty = False
    return repo


def merge_csv(repo: Repository, path: Path, kind: str) -> tuple[int, list[str]]:
    """Importa filas válidas en lote y conserva una sola acción para deshacer."""
    if kind not in ALL_FIELDS:
        raise DataError("Tipo de CSV desconocido")
    rows = _csv_read(Path(path), ALL_FIELDS[kind])
    before = repo.snapshot()
    accepted, errors = 0, []
    for number, row in enumerate(rows, start=2):
        try:
            repo.create(kind, row, remember=False)
            accepted += 1
        except DataError as exc:
            errors.append(f"Fila {number}: {exc}")
    if accepted:
        repo.history.push(before)
        repo.dirty = True
    return accepted, errors


def preview_csv_merge(repo: Repository, path: Path, kind: str) -> tuple[int, int, list[str]]:
    """Comprueba en una copia todos los aciertos/rechazos antes de mutar datos."""
    if kind not in ALL_FIELDS:
        raise DataError("Tipo de CSV desconocido")
    rows = _csv_read(Path(path), ALL_FIELDS[kind])
    staged = Repository()
    staged._restore(repo.snapshot())
    accepted, errors = 0, []
    for number, row in enumerate(rows, start=2):
        try:
            staged.create(kind, row, remember=False)
            accepted += 1
        except DataError as exc:
            errors.append(f"Fila {number}: {exc}")
    return len(rows), accepted, errors


def bundled_windows_cpp(root: Path) -> Path | None:
    """Usa el binario incluido solo si coincide con las fuentes y con su huella propia."""
    binary = root / "bin" / "windows" / "pea_cpp.exe"
    fingerprint = binary.with_suffix(".source-sha256")
    integrity = root / "bin" / "windows" / "pea_cpp.exe.sha256"
    if not binary.is_file() or not fingerprint.is_file() or not integrity.is_file():
        return None
    try:
        digest = hashlib.sha256()
        for source in (root / "CMakeLists.txt", root / "src/cpp/Taller2_REMR.cpp"):
            digest.update(source.relative_to(root).as_posix().encode("utf-8"))
            digest.update(b"\0")
            digest.update(source.read_bytes().replace(b"\r\n", b"\n"))
        if fingerprint.read_text(encoding="ascii").strip() != digest.hexdigest():
            return None
        binary_digest = hashlib.sha256(binary.read_bytes()).hexdigest()
        if integrity.read_text(encoding="ascii").strip() != binary_digest:
            return None
        return binary
    except (OSError, UnicodeError):
        pass
    return None


def cpp_executable(explicit: Path | None = None) -> Path:
    """Localiza o compila el ejecutable que atiende a la interfaz Tkinter."""
    root = Path(__file__).resolve().parents[2]
    if explicit is not None:
        binary = Path(explicit).expanduser().resolve()
        if not binary.is_file():
            raise DataError(f"No existe el backend C++: {binary}")
        return binary
    candidates = [root / "build" / "mingw" / "pea_cpp.exe",
                  root / "build" / "default" / "Release" / "pea_cpp.exe",
                  root / "build" / "default" / "pea_cpp.exe",
                  root / "build" / "Release" / "pea_cpp.exe",
                  root / "build" / "pea_cpp.exe",
                  root / "build" / "pea_cpp"]
    available = [item for item in candidates if item.is_file()]
    binary = max(available, key=lambda item: item.stat().st_mtime) if available else None
    newest_source = max((root / "CMakeLists.txt").stat().st_mtime,
                        (root / "src" / "cpp" / "Taller2_REMR.cpp").stat().st_mtime)
    if binary is not None and binary.stat().st_mtime >= newest_source:
        return binary
    if os.name == "nt":
        bundled = bundled_windows_cpp(root)
        if bundled is not None:
            return bundled
    if binary is None or binary.stat().st_mtime < newest_source:
        if shutil.which("cmake") is None:
            raise DataError("CMake no está instalado. Instálelo para compilar el backend C++.")
        mingw = bool(os.name == "nt" and shutil.which("g++") and shutil.which("mingw32-make"))
        build_dir = root / "build" / ("mingw" if mingw else "default")
        configure = ["cmake", "-S", str(root), "-B", str(build_dir)]
        if mingw:
            configure += ["-G", "MinGW Makefiles"]
        for command in (configure, ["cmake", "--build", str(build_dir), "--config", "Release"]):
            try:
                completed = subprocess.run(command, cwd=root, text=True, encoding="utf-8",
                                           stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                           check=False)
            except OSError as exc:
                raise DataError(f"No se pudo iniciar CMake: {exc}") from exc
            if completed.returncode:
                raise DataError("No se pudo compilar el backend C++:\n" + completed.stdout[-3000:])
        available = [item for item in candidates if item.is_file()]
        binary = max(available, key=lambda item: item.stat().st_mtime) if available else None
    if binary is None:
        raise DataError("CMake terminó sin generar pea_cpp; revise la carpeta build.")
    return binary


class CppRepository:
    """Adaptador local: Tkinter envía operaciones y el proceso C++ conserva el estado."""

    def __init__(self, binary: Path | None = None) -> None:
        executable = cpp_executable(binary)
        root = Path(__file__).resolve().parents[2]
        try:
            self.process = subprocess.Popen(
                [str(executable), "--api"], cwd=root,
                stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                text=True, encoding="utf-8", bufsize=1,
            )
        except OSError as exc:
            raise DataError(f"No se pudo iniciar el backend C++: {exc}") from exc
        self.dirty = False
        self._history_size = 0
        self._queue_size = 0
        try:
            hello = self._request("ping")
            if hello != {"backend": "cpp", "protocol": "1"}:
                raise DataError("El backend C++ usa un protocolo incompatible")
        except Exception:
            self.close()
            raise

    def _request(self, action: str, **fields: str) -> Any:
        if self.process.poll() is not None:
            detail = self.process.stderr.read().strip() if self.process.stderr else ""
            raise DataError("El backend C++ se cerró" + (f": {detail}" if detail else ""))
        assert self.process.stdin is not None and self.process.stdout is not None
        request = {"action": action, **fields}
        try:
            self.process.stdin.write(json.dumps(request, ensure_ascii=False, separators=(",", ":")) + "\n")
            self.process.stdin.flush()
            line = self.process.stdout.readline()
        except (OSError, UnicodeError) as exc:
            raise DataError(f"No se pudo comunicar con el backend C++: {exc}") from exc
        if not line:
            detail = self.process.stderr.read().strip() if self.process.stderr else ""
            raise DataError("El backend C++ no respondió" + (f": {detail}" if detail else ""))
        try:
            response = json.loads(line)
            state = response["state"]
            self.dirty = bool(state["dirty"])
            self._history_size = int(state["history_size"])
            self._queue_size = int(state["queue_size"])
            if not response["ok"]:
                raise DataError(response["error"])
            return response["result"]
        except (KeyError, TypeError, json.JSONDecodeError) as exc:
            raise DataError("Respuesta inválida del backend C++") from exc

    @staticmethod
    def _keys(key: str | tuple[str, str]) -> dict[str, str]:
        return {"key": key[0], "second": key[1]} if isinstance(key, tuple) else {"key": key, "second": ""}

    def rows(self, kind: str) -> list[dict[str, str]]:
        return self._request("rows", kind=kind)

    def page(self, kind: str, offset: int = 0, limit: int = PAGE_SIZE,
             query: str = "") -> dict[str, Any]:
        return self._request("page", kind=kind, offset=str(offset), limit=str(limit), query=query)

    def summary(self) -> dict[str, Any]:
        return self._request("summary")

    def get(self, kind: str, key: str | tuple[str, str]) -> dict[str, str] | None:
        return self._request("get", kind=kind, **self._keys(key))

    def related(self, kind: str, key: str, side: str) -> list[dict[str, str]]:
        return self._request("related", kind=kind, key=key, side=side)

    def related_page(self, kind: str, key: str, side: str, offset: int = 0,
                     limit: int = PAGE_SIZE, active_only: bool = False) -> dict[str, Any]:
        return self._request("related_page", kind=kind, key=key, side=side,
                             offset=str(offset), limit=str(limit), active="1" if active_only else "0")

    def create(self, kind: str, values: dict[str, str]) -> dict[str, str]:
        return self._request("create", kind=kind, values=json.dumps(values, ensure_ascii=False))

    def update(self, kind: str, key: str | tuple[str, str], values: dict[str, str]) -> dict[str, str]:
        return self._request("update", kind=kind, values=json.dumps(values, ensure_ascii=False),
                             **self._keys(key))

    def delete(self, kind: str, key: str | tuple[str, str]) -> None:
        self._request("delete", kind=kind, **self._keys(key))

    def toggle(self, kind: str, key: str | tuple[str, str]) -> None:
        self._request("toggle", kind=kind, **self._keys(key))

    def statistics(self, view: str = "Todos", selected_id: str = "", start: int | None = None,
                   end: int | None = None, category: str = "", status: str = "",
                   offset: int = 0, limit: int = PAGE_SIZE) -> dict[str, Any]:
        return self._request("statistics", view=view, selected=selected_id,
                             start="" if start is None else str(start),
                             end="" if end is None else str(end),
                             category=category, status=status, offset=str(offset), limit=str(limit))

    def history_size(self) -> int:
        return self._history_size

    def suggest_id(self, kind: str) -> str:
        return self._request("new_id", kind=kind)

    def undo(self) -> bool:
        return self._request("undo")

    def clear_history(self) -> None:
        self._request("clear_history")

    def reset(self) -> None:
        self._request("reset")

    def load(self, directory: Path) -> None:
        self._request("load", path=str(Path(directory).resolve()))

    def save(self, directory: Path) -> None:
        self._request("save", path=str(Path(directory).resolve()))

    def preview_csv(self, path: Path, kind: str) -> tuple[int, int, list[str]]:
        result = self._request("preview_csv", path=str(Path(path).resolve()), kind=kind)
        return result["total"], result["accepted"], result["errors"]

    def import_csv(self, path: Path, kind: str) -> tuple[int, list[str]]:
        result = self._request("import_csv", path=str(Path(path).resolve()), kind=kind)
        return result["accepted"], result["errors"]

    def queue_rows(self) -> list[dict[str, str]]:
        return self._request("queue_rows")
    def queue_page(self, offset: int = 0, limit: int = PAGE_SIZE) -> dict[str, Any]:
        return self._request("queue_page", offset=str(offset), limit=str(limit))
    def queue_front(self) -> dict[str, str] | None:
        return self._request("queue_front")
    def queue_size(self) -> int:
        return self._queue_size
    def enqueue_review(self, product_id: str, reason: str) -> None:
        self._request("enqueue_review", product_id=product_id, reason=reason)
    def process_review(self, status: str, observation: str) -> dict[str, str]:
        return self._request("process_review", status=status, observation=observation)
    def discard_review(self) -> dict[str, str]:
        return self._request("discard_review")

    def close(self) -> None:
        if self.process.poll() is None:
            try:
                self._request("shutdown")
                self.process.wait(timeout=2)
            except (DataError, OSError, subprocess.TimeoutExpired):
                self.process.kill()
                self.process.wait()
        for stream in (self.process.stdin, self.process.stdout, self.process.stderr):
            if stream is not None:
                stream.close()


class VisibleText(HTMLParser):
    BLOCKS = {"tr", "td", "th", "h1", "h2", "h3", "h4", "p", "li", "br", "div", "section", "article", "title"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.suppressed = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self.suppressed += 1
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self.suppressed = max(0, self.suppressed - 1)
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.suppressed:
            self.parts.append(data)

    def lines(self) -> list[str]:
        return [re.sub(r"\s+", " ", part).strip() for part in "".join(self.parts).splitlines() if part.strip()]


class LinkedTableRows(HTMLParser):
    """Conserva celdas y enlaces de tablas HTML para distinguir nombres de códigos CvLAC."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[tuple[str, str]]] = []
        self._row: list[tuple[str, str]] | None = None
        self._cell: list[str] | None = None
        self._href = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
            self._href = ""
        elif tag == "a" and self._cell is not None:
            self._href = dict(attrs).get("href") or self._href

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None:
            if self._row is not None:
                self._row.append((re.sub(r"\s+", " ", "".join(self._cell)).strip(), self._href))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row:
                self.rows.append(self._row)
            self._row = None


class WebPreview(NamedTuple):
    url: str
    format: str
    title: str
    lines: tuple[str, ...]
    metadata: dict[str, str]
    suggested_kind: str | None = None
    suggested_row: dict[str, str] | None = None
    csv_bytes: bytes | None = None
    tables: tuple[tuple[tuple[str, ...], ...], ...] = ()
    related_products: tuple[dict[str, str], ...] = ()
    related_members: tuple[tuple[dict[str, str], dict[str, str]], ...] = ()
    related_plan: dict[str, str] | None = None
    related_authorships: tuple[dict[str, str], ...] = ()


class GrupLACProduct(NamedTuple):
    row: dict[str, str]
    authors: tuple[str, ...]
    section: str


class PageText(VisibleText):
    """Recoge texto visible y metadatos explícitos, sin interpretar el tema del sitio."""

    META_KEYS = {
        "description", "og:description", "og:title", "twitter:title", "author", "date", "article:published_time",
        "citation_title", "citation_doi", "citation_date", "citation_publication_date",
        "citation_author", "citation_journal_title", "dc.title", "dc.date",
    }

    def __init__(self) -> None:
        super().__init__()
        self.metadata: dict[str, str] = {}
        self.title_parts: list[str] = []
        self.heading_parts: list[str] = []
        self._title = False
        self._heading = False
        self._json_ld = False
        self._json_parts: list[str] = []
        self.schemas: list[dict[str, Any]] = []
        self.tables: list[tuple[tuple[str, ...], ...]] = []
        self._table_level = 0
        self._table_rows: list[tuple[str, ...]] = []
        self._table_row: list[str] | None = None
        self._table_cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        if tag == "meta":
            key = (attributes.get("name") or attributes.get("property") or "").casefold()
            value = (attributes.get("content") or "").strip()
            if key in self.META_KEYS and value and key not in self.metadata:
                self.metadata[key] = value[:1000]
        if tag == "title":
            self._title = True
        if tag == "h1" and not self.heading_parts:
            self._heading = True
        if tag == "script" and attributes.get("type", "").casefold() == "application/ld+json":
            self._json_ld = True
            self._json_parts = []
        if tag == "table":
            self._table_level += 1
            if self._table_level == 1:
                self._table_rows = []
        elif self._table_level == 1 and tag == "tr":
            self._table_row = []
        elif self._table_level == 1 and tag in ("td", "th") and self._table_row is not None:
            self._table_cell = []
        super().handle_starttag(tag, attrs)

    def handle_endtag(self, tag: str) -> None:
        if tag == "script" and self._json_ld:
            self._json_ld = False
            try:
                schema = json.loads("".join(self._json_parts))
            except (ValueError, TypeError, RecursionError):
                schema = None
            self._collect_schema(schema)
        if tag == "title":
            self._title = False
        if tag == "h1":
            self._heading = False
        if self._table_level == 1 and tag in ("td", "th") and self._table_cell is not None:
            if self._table_row is not None and len(self._table_row) < 30:
                self._table_row.append(re.sub(r"\s+", " ", "".join(self._table_cell)).strip()[:500])
            self._table_cell = None
        elif self._table_level == 1 and tag == "tr" and self._table_row is not None:
            if self._table_row and len(self._table_rows) < 100:
                self._table_rows.append(tuple(self._table_row))
            self._table_row = None
        if tag == "table" and self._table_level:
            if self._table_level == 1 and self._table_rows and len(self.tables) < 5:
                self.tables.append(tuple(self._table_rows))
            self._table_level -= 1
        super().handle_endtag(tag)

    def handle_data(self, data: str) -> None:
        if self._json_ld:
            self._json_parts.append(data)
        if self._title:
            self.title_parts.append(data)
        if self._heading:
            self.heading_parts.append(data)
        if self._table_cell is not None:
            self._table_cell.append(data)
        super().handle_data(data)

    def _collect_schema(self, item: Any, depth: int = 0) -> None:
        if depth > 64:
            return
        if isinstance(item, list):
            for value in item:
                self._collect_schema(value, depth + 1)
        elif isinstance(item, dict):
            if isinstance(item.get("@graph"), list):
                self._collect_schema(item["@graph"], depth + 1)
            if "@type" in item:
                self.schemas.append(item)


def _web_source(url: str) -> str:
    return f"{urlparse(url).hostname}, consultado {date.today().isoformat()}"


def _redact_url(url: str, keep_params: tuple[str, ...] = ("nro", "cod_rh")) -> str:
    """Elimina parámetros de consulta sensibles; conserva solo los necesarios para reabrir."""
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        return url
    params = parse_qs(parsed.query)
    kept = {key: params[key][0] for key in keep_params if key in params}
    query = urlencode(sorted(kept.items()))
    return parsed._replace(query=query, fragment="").geturl()


def _publication_year(value: str) -> str:
    match = re.search(r"(?<!\d)(?:19|20)\d{2}(?!\d)", value)
    if match and 1900 <= int(match.group()) <= date.today().year + 1:
        return match.group()
    return ""


def _is_scholarly_article(item: dict[str, Any]) -> bool:
    declared = item.get("@type")
    types = [declared] if isinstance(declared, str) else declared if isinstance(declared, list) else []
    return any(isinstance(value, str) and value.rsplit("/", 1)[-1] == "ScholarlyArticle" for value in types)


def parse_web_page(text: str, url: str, format: str = "HTML") -> WebPreview:
    """Previsualiza cualquier página de texto; solo propone campos con etiquetas fiables."""
    metadata: dict[str, str] = {}
    suggestion: tuple[str, dict[str, str]] | None = None
    tables: tuple[tuple[tuple[str, ...], ...], ...] = ()
    related_products: tuple[dict[str, str], ...] = ()
    related_members: tuple[tuple[dict[str, str], dict[str, str]], ...] = ()
    related_plan: dict[str, str] | None = None
    related_authorships: tuple[dict[str, str], ...] = ()
    if format == "HTML":
        parser = PageText()
        parser.feed(text)
        lines = parser.lines()
        metadata = parser.metadata.copy()
        tables = tuple(parser.tables)
        if tables:
            metadata["Tablas HTML"] = str(len(tables))
        title = re.sub(r"\s+", " ", "".join(parser.title_parts)).strip()
        heading = re.sub(r"\s+", " ", "".join(parser.heading_parts)).strip()
        title = title or metadata.get("og:title", "") or metadata.get("twitter:title", "") or heading or metadata.get("citation_title", "")
        if urlparse(url).hostname == "scienti.minciencias.gov.co":
            try:
                suggestion = parse_scienti_html(text, url)
                if suggestion[0] == "grupos":
                    members = parse_gruplac_members(text, url)
                    products = parse_gruplac_products(text, url)
                    related_members = tuple(members)
                    related_products = tuple(item.row for item in products)
                    related_plan = parse_gruplac_plan(text, url, suggestion[1]["objetivos"])
                    related_authorships = tuple(match_gruplac_authors(products, members))
            except DataError:
                pass
        if suggestion is None:
            article = next((item for item in parser.schemas if _is_scholarly_article(item)), None)
            article_title = metadata.get("citation_title", "")
            if article and not article_title and isinstance(article.get("headline") or article.get("name"), str):
                article_title = article.get("headline") or article.get("name")
            if article_title:
                published = metadata.get("citation_publication_date") or metadata.get("citation_date") or (
                    article.get("datePublished", "") if article else "")
                doi = metadata.get("citation_doi") or (article.get("doi", "") if article else "")
                if article:
                    for key, value in (("schema.org: título", article_title),
                                       ("schema.org: fecha", published), ("schema.org: DOI", doi)):
                        if value and key not in metadata:
                            metadata[key] = str(value)[:1000]
                row = {"id": new_id("P"), "titulo": article_title, "anio": _publication_year(str(published)),
                       "doi": str(doi) if isinstance(doi, str) else "", "url": _redact_url(url), "fuente": _web_source(url)}
                suggestion = ("productos", row)
        if not title:
            title = lines[0] if lines else urlparse(url).hostname or url
    elif format == "CSV":
        reader = csv.reader(io.StringIO(text))
        try:
            columns = next(reader)
        except StopIteration:
            columns = []
        metadata["Columnas CSV"] = ", ".join(columns[:30]) or "Archivo vacío"
        lines = [f"Encabezados: {metadata['Columnas CSV']}"]
        for index, row in enumerate(reader, start=1):
            if index > 100:
                lines.append("… Vista previa limitada a 100 filas")
                break
            lines.append(f"Fila {index}: " + " | ".join(row[:30]))
        title = "CSV: " + (urlparse(url).path.rsplit("/", 1)[-1] or urlparse(url).hostname or url)
    else:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        title = lines[0][:200] if lines else urlparse(url).path.rsplit("/", 1)[-1] or url
    if not lines and not metadata:
        lines = ["No se encontró texto legible en la respuesta. El sitio podría requerir JavaScript o acceso especial."]
    if suggestion:
        kind, row = suggestion
    else:
        kind, row = None, None
    return WebPreview(url, format, title[:250], tuple(lines[:300]), metadata, kind, row, None, tables,
                      related_products, related_members, related_plan, related_authorships)


def _resolve_public_addrs(host: str, port: int) -> list[str]:
    try:
        addresses = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    except (OSError, ValueError) as exc:
        raise DataError(f"No se pudo resolver el sitio: {exc}") from exc
    resolved: list[str] = []
    for item in addresses:
        address = ipaddress.ip_address(item[4][0])
        if not address.is_global:
            raise DataError("No se permiten direcciones locales o privadas")
        if item[4][0] not in resolved:
            resolved.append(item[4][0])
    if not resolved:
        raise DataError("No se pudo resolver el sitio")
    return resolved


def _validate_public_url(url: str) -> str:
    try:
        parsed = urlparse(url.strip())
        host = parsed.hostname
        port = parsed.port
    except ValueError as exc:
        raise DataError("URL inválida") from exc
    if parsed.scheme not in ("http", "https") or not host or parsed.username or parsed.password:
        raise DataError("Indique una URL pública HTTP o HTTPS, sin usuario ni contraseña")
    if host.casefold() == "localhost" or host.casefold().endswith((".localhost", ".local")):
        raise DataError("No se permiten direcciones locales o privadas")
    if port is not None and port not in (80, 443):
        raise DataError("Solo se permiten los puertos 80 y 443")
    try:
        literal_address = ipaddress.ip_address(host)
    except ValueError:
        literal_address = None
    if literal_address is not None and not literal_address.is_global:
        raise DataError("No se permiten direcciones locales o privadas")
    _resolve_public_addrs(host, port or (443 if parsed.scheme == "https" else 80))
    return url.strip()


class _PinnedHTTPConnection(http.client.HTTPConnection):
    def connect(self) -> None:
        address = _resolve_public_addrs(self.host, self.port)[0]
        self.sock = socket.create_connection((address, self.port), self.timeout, self.source_address)


class _PinnedHTTPSConnection(http.client.HTTPSConnection):
    def connect(self) -> None:
        address = _resolve_public_addrs(self.host, self.port)[0]
        self.sock = socket.create_connection((address, self.port), self.timeout, self.source_address)
        if self._tunnel_host:
            self._tunnel()
        self.sock = self._context.wrap_socket(self.sock, server_hostname=self._tunnel_host or self.host)


class _PinnedHTTPHandler(HTTPHandler):
    def http_open(self, req: Request) -> Any:
        return self.do_open(_PinnedHTTPConnection, req)


class _PinnedHTTPSHandler(HTTPSHandler):
    def https_open(self, req: Request) -> Any:
        return self.do_open(_PinnedHTTPSConnection, req, context=self._context)


class PublicRedirects(HTTPRedirectHandler):
    def redirect_request(self, req: Request, fp: Any, code: int, msg: str,
                         headers: Any, newurl: str) -> Request | None:
        _validate_public_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def _build_public_opener() -> Any:
    return build_opener(ProxyHandler({}), _PinnedHTTPHandler(),
                        _PinnedHTTPSHandler(), PublicRedirects())


def fetch_web_page(url: str) -> WebPreview:
    """Descarga una sola URL pública y previsualiza HTML, texto, CSV o PDF de texto."""
    url = _validate_public_url(url)
    request = Request(url, headers={"User-Agent": "PEA-i-UPC/1.0 (academic project)",
                                    "Accept": "text/html,text/plain,text/csv,application/pdf;q=0.9,*/*;q=0.1"})
    try:
        opener = _build_public_opener()
        with opener.open(request, timeout=15) as response:
            final_url = _validate_public_url(response.geturl())
            content_type = response.headers.get_content_type().casefold()
            charset = response.headers.get_content_charset()
            content = response.read(8_000_001)
    except (OSError, TimeoutError) as exc:
        raise DataError(f"No se pudo descargar la URL: {exc}") from exc
    if len(content) > 8_000_000:
        raise DataError("La respuesta supera 8 MB; use un archivo CSV local")
    if content.startswith(b"%PDF") or content_type == "application/pdf":
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "pagina.pdf"
            path.write_bytes(content)
            preview = parse_web_page(extract_pdf_text(path), final_url, "PDF de texto")
            if urlparse(final_url).hostname == "scienti.minciencias.gov.co":
                try:
                    kind, row = parse_scienti_pdf(path, final_url)
                    preview = preview._replace(suggested_kind=kind, suggested_row=row)
                except DataError:
                    pass
            return preview
    csv_file = content_type in ("text/csv", "application/csv", "application/vnd.ms-excel") or urlparse(final_url).path.lower().endswith(".csv")
    html_file = content_type in ("text/html", "application/xhtml+xml") or content[:500].lstrip().lower().startswith((b"<!doctype html", b"<html"))
    if not (csv_file or html_file or content_type.startswith("text/")):
        raise DataError(f"Formato web no legible: {content_type}")
    if not charset and html_file:
        match = re.search(rb"charset\s*=\s*['\"]?([A-Za-z0-9_-]+)", content[:5000], re.IGNORECASE)
        charset = match.group(1).decode("ascii") if match else None
    try:
        decoded = content.decode(charset or "utf-8-sig", errors="replace")
    except LookupError as exc:
        raise DataError(f"Codificación desconocida: {charset}") from exc
    preview = parse_web_page(decoded, final_url, "CSV" if csv_file else "HTML" if html_file else "Texto")
    return preview._replace(csv_bytes=content) if csv_file else preview


def chart_values(values: dict[str, int]) -> dict[str, int]:
    """Los campos sin valor cuentan en el total, pero no forman una barra de datos."""
    return {key: value for key, value in values.items() if key.strip().casefold() != "sin dato" and value > 0}


_GRUPLAC_LEVELS = ("A1", "A", "B", "C", "D", "Reconocido")


def _gruplac_category(value: str) -> str:
    """Reduce el texto de 'Clasificación' a un nivel válido (A1..D, Reconocido) o vacío."""
    token = value.split(" ")[0].strip()
    return token if token in _GRUPLAC_LEVELS else ""


def parse_scienti_html(html: str, url: str) -> tuple[str, dict[str, str]]:
    """Extrae campos identificables; evita atribuir productos a un perfil ambiguo."""
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != "scienti.minciencias.gov.co":
        raise DataError("Solo se admiten URL públicas HTTPS de Scienti Minciencias")
    parser = VisibleText()
    parser.feed(html)
    lines = parser.lines()
    if not lines:
        raise DataError("La página no contiene texto legible")
    code = parse_qs(parsed.query)
    if "/gruplac/" in parsed.path and code.get("nro"):
        candidates = [line for line in lines if line.lower() not in ("gruplac - plataforma scienti - colombia", "datos básicos")]
        name = next((line for line in candidates if line in lines[:15] and len(line) > 2 and not line.lower().startswith(("inicio", "menú", "buscar"))), "")
        if not name:
            raise DataError("No se identificó el nombre del grupo")
        text = "\n".join(lines)
        def after(label: str) -> str:
            match = re.search(rf"(?im)^{re.escape(label)}\s*(?:[:|])?\s*(.+)$", text)
            if match:
                return match.group(1).strip()
            try:
                index = next(i for i, line in enumerate(lines) if line.casefold() == label.casefold())
                return lines[index + 1] if index + 1 < len(lines) else ""
            except StopIteration:
                return ""
        def block(label: str, stop: str) -> str:
            try:
                first = next(i for i, line in enumerate(lines) if line.casefold().startswith(label.casefold()))
            except StopIteration:
                return ""
            collected = [lines[first].split(":", 1)[1].strip()] if ":" in lines[first] else []
            for line in lines[first + 1:]:
                if line.casefold().startswith(stop.casefold()):
                    break
                collected.append(line)
            return "\n".join(part for part in collected if part)
        formed = after("Año y mes de formación")
        group = {
            "id": f"G-{code['nro'][0]}", "nombre": name,
            "codigo_gruplac": code["nro"][0], "responsable": after("Líder"),
            "categoria": _gruplac_category(after("Clasificación")),
            "descripcion": block("Estado del arte", "Objetivos"),
            "objetivos": block("Objetivos", "Retos"),
            "vision": after("Visión"),
            "lineas": block("Líneas de investigación declaradas por el grupo", "Integrantes del grupo"),
            "url": _redact_url(url), "fuente": f"Scienti, consultado {date.today().isoformat()}; formación: {formed}",
        }
        return "grupos", clean_row("grupos", group)
    if "/cvlac/" in parsed.path and code.get("cod_rh"):
        text = "\n".join(lines)
        match = re.search(r"(?im)^Nombre\s+([^\n]+)$", text)
        if not match:
            raise DataError("No se identificó el nombre del investigador")
        name = match.group(1).strip()
        category = ""
        for index, line in enumerate(lines):
            if line.casefold() == "categoría" and index + 1 < len(lines):
                category = lines[index + 1].split(" con vigencia")[0].strip()
                break
        if not category:
            found = re.search(r"Categoría Investigador\s+([^\n]+)", text, re.IGNORECASE)
            category = found.group(1).split(" con vigencia")[0].strip() if found else ""
        researcher = {
            "id": f"I-{code['cod_rh'][0]}", "nombre": name,
            "codigo_cvlac": code["cod_rh"][0],
            "categoria": category,
            "url": _redact_url(url), "fuente": f"Scienti, consultado {date.today().isoformat()}",
        }
        return "investigadores", clean_row("investigadores", researcher)
    raise DataError("La URL debe ser una ficha GrupLAC o CvLAC con su código")


def _require_gruplac_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != "scienti.minciencias.gov.co" or "/gruplac/" not in parsed.path:
        raise DataError("La fuente debe ser una ficha GrupLAC HTTPS")
    code = parse_qs(parsed.query).get("nro", [""])[0]
    if not code:
        raise DataError("La ficha GrupLAC no tiene código de grupo")
    return code


def _name_key(value: str) -> str:
    plain = "".join(char for char in unicodedata.normalize("NFKD", value) if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", plain).casefold().strip()


def parse_gruplac_members(html: str, url: str) -> list[tuple[dict[str, str], dict[str, str]]]:
    """Importa el censo explícito; conserva meses textuales sin inventar un día ISO."""
    code = _require_gruplac_url(url)
    parser = LinkedTableRows()
    parser.feed(html)
    first = next((i + 1 for i, row in enumerate(parser.rows)
                  if len(row) >= 4 and _name_key(row[0][0]) == "nombre"
                  and _name_key(row[1][0]) == "vinculacion"
                  and "inicio" in _name_key(row[3][0])), None)
    if first is None:
        return []
    members: list[tuple[dict[str, str], dict[str, str]]] = []
    seen: set[str] = set()
    for row in parser.rows[first:]:
        if len(row) < 4:
            if members:
                break
            continue
        match = re.match(r"^\d+\.-\s*(.+)$", row[0][0])
        if not match:
            if members:
                break
            continue
        name = match.group(1).strip()
        href = row[0][1]
        profile = urlparse(href)
        profile_code = ""
        if profile.scheme == "https" and profile.hostname == "scienti.minciencias.gov.co" and "/cvlac/" in profile.path:
            profile_code = parse_qs(profile.query).get("cod_rh", [""])[0]
        person_id = ("I-" + profile_code) if profile_code else (
            "I-GR-" + hashlib.sha1((code + "|" + _name_key(name)).encode("utf-8")).hexdigest()[:12].upper())
        if person_id in seen:
            continue
        seen.add(person_id)
        period = row[3][0].strip()
        role = row[1][0].strip() or "Integrante"
        ended = bool(re.search(r"\s-\s(?:19|20)\d{2}/\d{1,2}\s*$", period))
        person = clean_row("investigadores", {
            "id": person_id, "nombre": name, "codigo_cvlac": profile_code,
            "url": _redact_url(href) if profile_code else _redact_url(url),
            "fuente": f"GrupLAC, censo de integrantes, consultado {date.today().isoformat()}; vinculación: {period}",
        })
        membership = clean_row("membresias", {
            "grupo_id": "G-" + code, "investigador_id": person_id,
            "rol": f"{role} · {period}" if period else role,
            "activo": "0" if ended else "1",
        })
        members.append((person, membership))
    return members


def parse_gruplac_plan(html: str, url: str, objectives: str = "") -> dict[str, str] | None:
    code = _require_gruplac_url(url)
    parser = VisibleText()
    parser.feed(html)
    lines = parser.lines()
    try:
        first = lines.index("Plan Estratégico") + 1
    except ValueError:
        return None
    last = next((i for i in range(first, len(lines)) if lines[i].startswith("Estado del arte")), first)
    activities = "\n".join(lines[first:last]).strip()
    if not activities and not objectives:
        return None
    return clean_row("planes", {
        "id": "PL-" + code, "grupo_id": "G-" + code, "nombre": "Plan Estratégico",
        "objetivo": objectives, "actividad": activities,
    })


GRUPLAC_SECTIONS = (
    ("Artículos publicados", "Libros publicados", "Producción bibliográfica"),
    ("Libros publicados", "Capítulos de libro publicados", "Producción bibliográfica"),
    ("Capítulos de libro publicados", "Documentos de trabajo", "Producción bibliográfica"),
    ("Documentos de trabajo", "Otra publicación divulgativa", "Producción bibliográfica"),
    ("Otra publicación divulgativa", "Otros artículos publicados", "Producción bibliográfica"),
    ("Otros artículos publicados", "Libros de formación", "Producción bibliográfica"),
    ("Libros de formación", "Libros de divulgación y/o Compilación de divulgación", "Producción bibliográfica"),
    ("Otros Libros publicados", "Traducciones", "Producción bibliográfica"),
    ("Otros productos tecnológicos", "Prototipos", "Producción técnica y tecnológica"),
    ("Softwares", "Empresas de base tecnológica", "Producción técnica y tecnológica"),
    ("Empresas de base tecnológica", "APROPIACIÓN SOCIAL Y DIVULGACIÓN PÚBLICA DE LA CIENCIA", "Producción técnica y tecnológica"),
)


def parse_gruplac_products(html: str, url: str) -> list[GrupLACProduct]:
    """Extrae secciones con tipo, título y año explícitos; une fichas repetidas por DOI o título."""
    _require_gruplac_url(url)
    parser = VisibleText()
    parser.feed(html)
    lines = parser.lines()
    by_identity: dict[str, GrupLACProduct] = {}
    for section, stop, family in GRUPLAC_SECTIONS:
        try:
            first = lines.index(section) + 1
        except ValueError:
            continue
        last = next((i for i in range(first, len(lines)) if lines[i] == stop), len(lines))
        for index in range(first, last):
            match = re.match(r"^\d+\.-\s*([^:]+?)\s*:\s*(.+)$", lines[index])
            if not match:
                continue
            kind, title = match.group(1).strip(), match.group(2).strip()
            if not title:
                continue
            details: list[str] = []
            authors: list[str] = []
            for line in lines[index + 1:min(index + 16, last)]:
                if re.match(r"^\d+\.-\s", line):
                    break
                if line.startswith("Autores:"):
                    authors.extend(name.strip() for name in line[8:].split(",") if name.strip())
                    break
                details.append(line)
            detail = " ".join(details)
            year_match = re.search(r"(?<![\d-])((?:19|20)\d{2})(?=\s*(?:[,.;]|vol:|$))", detail)
            if not year_match or not _publication_year(year_match.group(1)):
                continue
            year = year_match.group(1)
            doi_match = re.search(r"\bDOI:\s*(10\.\S+)", detail, re.IGNORECASE)
            doi = doi_match.group(1).rstrip(".,;)") if doi_match else ""
            identity = "doi:" + doi.casefold() if doi else "title:" + _name_key(title) + "|" + year
            if not doi and section != "Artículos publicados":
                identity += "|" + _name_key(kind)
            previous = by_identity.get(identity)
            if previous:
                combined = list(previous.authors)
                known = {_name_key(name) for name in combined}
                for name in authors:
                    if _name_key(name) not in known:
                        combined.append(name)
                        known.add(_name_key(name))
                by_identity[identity] = previous._replace(authors=tuple(combined))
                continue
            stable_id = "P-GR-" + hashlib.sha1(identity.encode("utf-8")).hexdigest()[:12].upper()
            row = clean_row("productos", {
                "id": stable_id, "titulo": title, "anio": year,
                "familia": family, "tipologia": kind,
                "doi": doi, "url": _redact_url(url),
                "fuente": f"GrupLAC, sección {section}, consultado {date.today().isoformat()}",
            })
            by_identity[identity] = GrupLACProduct(row, tuple(authors), section)
    return list(by_identity.values())


def parse_gruplac_articles(html: str, url: str) -> list[dict[str, str]]:
    """Compatibilidad: artículos de revista de la sección propia de GrupLAC."""
    return [entry.row for entry in parse_gruplac_products(html, url) if entry.section == "Artículos publicados"]


def match_gruplac_authors(products: list[GrupLACProduct],
                          members: list[tuple[dict[str, str], dict[str, str]]]) -> list[dict[str, str]]:
    by_name = {_name_key(person["nombre"]): person["id"] for person, _ in members}
    relations: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for product in products:
        for author in product.authors:
            researcher_id = by_name.get(_name_key(author))
            pair = (product.row["id"], researcher_id or "")
            if not researcher_id or pair in seen:
                continue
            seen.add(pair)
            relations.append(clean_row("autorias", {"producto_id": pair[0], "investigador_id": pair[1],
                                                   "rol": "Autor listado en GrupLAC"}))
    return relations


def import_gruplac_preview(repository: Repository | CppRepository, preview: WebPreview,
                           reviewed_group: dict[str, str] | None = None) -> dict[str, Any]:
    """Incorpora datos respaldados por una ficha, respetando registros ya existentes."""
    if preview.suggested_kind != "grupos" or not preview.suggested_row:
        raise DataError("La vista previa no contiene una ficha de grupo")
    source_group = preview.suggested_row
    group = clean_row("grupos", reviewed_group or source_group)
    if group["codigo_gruplac"] != source_group["codigo_gruplac"]:
        raise DataError("El código GrupLAC revisado no coincide con la fuente")
    group_id = group["id"]
    counts: dict[str, Any] = {kind: 0 for kind in (
        "grupos", "investigadores", "membresias", "planes", "productos", "grupos_productos", "autorias")}
    counts["existentes"] = 0
    counts["errors"] = []

    def add(kind: str, row: dict[str, str], key: str | tuple[str, str]) -> bool:
        existing = repository.get(kind, key)
        if existing is not None:
            if kind == "grupos" and existing["codigo_gruplac"] != source_group["codigo_gruplac"]:
                counts["errors"].append(f"{kind} {key}: ID ocupado por otro grupo")
                return False
            if kind == "investigadores" and existing["codigo_cvlac"] != row["codigo_cvlac"]:
                counts["errors"].append(f"{kind} {key}: ID ocupado por otro perfil")
                return False
            if kind == "productos" and (existing["titulo"] != row["titulo"] or existing["anio"] != row["anio"]):
                counts["errors"].append(f"{kind} {key}: ID ocupado por otro producto")
                return False
            counts["existentes"] += 1
            return True
        try:
            repository.create(kind, row)
        except (DataError, OSError, ValueError) as exc:
            counts["errors"].append(f"{kind} {key}: {exc}")
            return False
        counts[kind] += 1
        return True

    if not add("grupos", group, group_id):
        return counts
    available_people: set[str] = set()
    for person, _ in preview.related_members:
        if add("investigadores", person, person["id"]):
            available_people.add(person["id"])
    if preview.related_plan:
        plan = {**preview.related_plan, "grupo_id": group_id}
        add("planes", plan, plan["id"])
    for person, membership in preview.related_members:
        if person["id"] in available_people:
            row = {**membership, "grupo_id": group_id}
            add("membresias", row, (group_id, person["id"]))
    available_products: set[str] = set()
    for product in preview.related_products:
        if add("productos", product, product["id"]):
            available_products.add(product["id"])
    for product in preview.related_products:
        if product["id"] in available_products:
            row = {"grupo_id": group_id, "producto_id": product["id"],
                   "origen": "GrupLAC: producto listado en la ficha"}
            add("grupos_productos", row, (group_id, product["id"]))
    for authorship in preview.related_authorships:
        product_id, researcher_id = authorship["producto_id"], authorship["investigador_id"]
        if product_id in available_products and researcher_id in available_people:
            add("autorias", authorship, (product_id, researcher_id))
    return counts


def extract_docx_html(path: Path) -> str:
    """Convierte un .docx a fragmento HTML con mammoth (instalación opcional)."""
    _check_local_file(path, archive=True)
    try:
        import mammoth
    except ImportError as exc:
        raise DataError("Para importar Word instale mammoth (pip install mammoth)") from exc
    try:
        with path.open("rb") as stream:
            html = mammoth.convert_to_html(stream).value
            if len(html) > MAX_PDF_TEXT:
                raise DataError("El documento Word produce demasiado texto")
            return html
    except DataError:
        raise
    except Exception as exc:
        raise DataError(f"No se pudo leer el documento Word: {exc}") from exc


def extract_pdf_text(path: Path) -> str:
    """Admite PDF de texto si pdftotext (Poppler) está instalado."""
    _check_local_file(path)
    if shutil.which("pdftotext") is None:
        raise DataError("Para importar PDF instale Poppler (pdftotext) o use CSV")

    def limit_resources() -> None:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (MAX_PDF_MEMORY, MAX_PDF_MEMORY))
        resource.setrlimit(resource.RLIMIT_CPU, (PDF_TIMEOUT, PDF_TIMEOUT))

    try:
        process = subprocess.Popen(
            ["pdftotext", "-layout", str(path), "-"],
            stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
            start_new_session=(os.name != "nt"),
            preexec_fn=limit_resources if os.name != "nt" else None,
        )
    except (OSError, ValueError) as exc:
        raise DataError(f"No se pudo leer el PDF: {exc}") from exc
    chunks: list[str] = []
    failures: list[Exception] = []

    def collect() -> None:
        total = 0
        try:
            assert process.stdout is not None
            while chunk := process.stdout.read(65536):
                total += len(chunk)
                if total > MAX_PDF_TEXT:
                    raise DataError("El PDF produce demasiado texto")
                chunks.append(chunk)
        except (OSError, UnicodeError, DataError) as exc:
            failures.append(exc)

    reader = threading.Thread(target=collect, daemon=True)
    reader.start()
    reader.join(PDF_TIMEOUT)
    if reader.is_alive() or failures:
        process.kill()
        process.wait()
        reader.join(timeout=1)
        if process.stdout is not None:
            process.stdout.close()
        if reader.is_alive():
            raise DataError("Se agotó el tiempo al leer el PDF")
        if failures:
            raise DataError(f"No se pudo leer el PDF: {failures[0]}") from failures[0]
        raise DataError("Se agotó el tiempo al leer el PDF")
    try:
        process.wait(timeout=2)
    except subprocess.TimeoutExpired as exc:
        process.kill()
        process.wait()
        raise DataError("Se agotó el tiempo al leer el PDF") from exc
    finally:
        if process.stdout is not None:
            process.stdout.close()
    if process.returncode != 0:
        raise DataError("No se pudo leer el PDF")
    text = "".join(chunks)
    if not text.strip():
        raise DataError("El PDF no contiene texto seleccionable; requiere OCR")
    return text


def parse_scienti_pdf(path: Path, source_url: str) -> tuple[str, dict[str, str]]:
    """Procesa PDFs impresos de GrupLAC/CvLAC con URL original conocida."""
    text = extract_pdf_text(path)
    # El exportador PDF presenta texto en bloques; reutilizamos la misma extracción
    # de campos, con saltos de línea preservados para no inventar asociaciones.
    from html import escape
    html = "<p>" + "</p><p>".join(escape(line) for line in text.splitlines()) + "</p>"
    return parse_scienti_html(html, source_url)


def new_id(prefix: str) -> str:
    return prefix + "-" + uuid.uuid4().hex[:8].upper()


def open_gui_repository(*, python_backend: bool = False,
                        cpp_binary: Path | None = None) -> tuple[Repository | CppRepository, str]:
    """Abre el motor disponible y devuelve un aviso si se usó el respaldo Python."""
    if python_backend:
        return Repository(), ""
    try:
        return CppRepository(cpp_binary), ""
    except DataError as exc:
        if cpp_binary is not None:
            raise
        return Repository(), str(exc).splitlines()[0]


def run_gui(data_dir: Path | None = None, *, python_backend: bool = False,
            cpp_binary: Path | None = None) -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, simpledialog, ttk

    class RecordDialog(tk.Toplevel):
        def __init__(self, parent: tk.Misc, kind: str, initial: dict[str, str] | None = None,
                     choices: dict[str, list[str]] | None = None, suggested_id: str = "",
                     external: bool = False):
            super().__init__(parent)
            self.title(("Revisar " if external else "Editar " if initial else "Crear ") + kind.replace("_", " "))
            self.transient(parent)
            self.resizable(True, True)
            self.geometry("760x650")
            self.result: dict[str, str] | None = None
            self.kind = kind
            self.inputs: dict[str, Any] = {}
            self.columnconfigure(0, weight=1)
            self.rowconfigure(0, weight=1)
            canvas = tk.Canvas(self, highlightthickness=0)
            scrollbar = ttk.Scrollbar(self, orient="vertical", command=canvas.yview)
            canvas.configure(yscrollcommand=scrollbar.set)
            canvas.grid(row=0, column=0, sticky="nsew")
            scrollbar.grid(row=0, column=1, sticky="ns")
            body = ttk.Frame(canvas, padding=12)
            window = canvas.create_window((0, 0), window=body, anchor="nw")
            body.bind("<Configure>", lambda _event: canvas.configure(scrollregion=canvas.bbox("all")))
            canvas.bind("<Configure>", lambda event: canvas.itemconfigure(window, width=event.width))
            body.columnconfigure(1, weight=1)
            fields = ALL_FIELDS[kind]
            for index, field in enumerate(fields):
                ttk.Label(body, text=LABELS.get(field, field)).grid(row=index, column=0, sticky="nw", padx=4, pady=3)
                value = initial.get(field, "") if initial else ""
                if field == "activo" and not initial:
                    value = "1"
                if field == "id" and not initial:
                    value = suggested_id
                if field in LONG_FIELDS:
                    widget = tk.Text(body, height=2, width=58, wrap="word")
                    widget.insert("1.0", value)
                elif field in ("activo", "validacion"):
                    options = ("1", "0") if field == "activo" else ("pendiente", "validado", "rechazado")
                    widget = ttk.Combobox(body, values=options, state="readonly", width=54)
                    widget.set(value)
                elif choices and field in choices:
                    widget = ttk.Combobox(body, values=choices[field], width=54)
                    widget.set(value)
                else:
                    widget = ttk.Entry(body, width=57)
                    widget.insert(0, value)
                widget.grid(row=index, column=1, sticky="ew", padx=4, pady=3)
                if initial and not external and (field == "id" or field in RELATION_ENDS.get(kind, ())[:2]):
                    widget.configure(state="disabled")
                self.inputs[field] = widget
            actions = ttk.Frame(body)
            actions.grid(row=len(fields), column=0, columnspan=2, pady=12)
            ttk.Button(actions, text="Guardar", command=self.accept).pack(side="left", padx=5)
            ttk.Button(actions, text="Cancelar", command=self.destroy).pack(side="left", padx=5)
            self.bind("<Escape>", lambda _event: self.destroy())
            self.grab_set()
            self.wait_window()

        def accept(self) -> None:
            result = {}
            for field, widget in self.inputs.items():
                if isinstance(widget, tk.Text):
                    result[field] = widget.get("1.0", "end-1c").strip()
                else:
                    result[field] = widget.get().strip()
            self.result = result
            self.destroy()

    class EntityTab(ttk.Frame):
        def __init__(self, parent: tk.Misc, app: "App", kind: str):
            super().__init__(parent, padding=8)
            self.app, self.kind = app, kind
            self.search = tk.StringVar()
            self.offset = 0
            self._search_timer: str | None = None
            toolbar = ttk.Frame(self)
            toolbar.pack(fill="x")
            ttk.Label(toolbar, text="Buscar").pack(side="left")
            ttk.Entry(toolbar, textvariable=self.search, width=26).pack(side="left", padx=5)
            self.search.trace_add("write", lambda *_: self._search_changed())
            for label, action in (("Nuevo", self.create), ("Editar", self.edit),
                                  ("Activar/desactivar", self.toggle), ("Eliminar", self.delete),
                                  ("Información", self.open_detail)):
                ttk.Button(toolbar, text=label, command=action).pack(side="left", padx=3)
            visible = {
                "grupos": ("id", "nombre", "categoria", "activo"),
                "investigadores": ("id", "nombre", "afiliacion", "categoria", "activo"),
                "productos": ("id", "titulo", "anio", "tipologia", "categoria", "validacion", "activo"),
                "planes": ("id", "grupo_id", "nombre", "inicio", "fin", "activo"),
            }[kind]
            table_frame = ttk.Frame(self)
            table_frame.pack(fill="both", expand=True, pady=(7, 0))
            table_frame.columnconfigure(0, weight=1)
            table_frame.rowconfigure(0, weight=1)
            self.table = ttk.Treeview(table_frame, columns=visible, show="headings", selectmode="browse", height=14)
            self.table.tag_configure("stripe", background=self.app.colors["stripe"])
            for field in visible:
                self.table.heading(field, text=LABELS[field])
                self.table.column(field, width=85 if field in ("id", "anio", "activo") else 180, stretch=True)
            table_y = ttk.Scrollbar(table_frame, orient="vertical", command=self.table.yview)
            table_x = ttk.Scrollbar(table_frame, orient="horizontal", command=self.table.xview)
            self.table.configure(yscrollcommand=table_y.set, xscrollcommand=table_x.set)
            self.table.grid(row=0, column=0, sticky="nsew")
            table_y.grid(row=0, column=1, sticky="ns")
            table_x.grid(row=1, column=0, sticky="ew")
            self.table.bind("<Double-1>", lambda _event: self.open_detail())
            pager = ttk.Frame(self)
            pager.pack(fill="x")
            ttk.Button(pager, text="Anterior", command=lambda: self._move(-1)).pack(side="left")
            self.page_label = tk.StringVar()
            ttk.Label(pager, textvariable=self.page_label).pack(side="left", padx=10)
            ttk.Button(pager, text="Siguiente", command=lambda: self._move(1)).pack(side="left")
            if kind in ("grupos", "investigadores"):
                ttk.Button(pager, text="Consultar fuente", command=self.open_source).pack(side="right")
            self.visible = visible

        def selected(self) -> str | None:
            selection = self.table.selection()
            return selection[0] if selection else None

        def _search_changed(self) -> None:
            self.offset = 0
            if self._search_timer is not None:
                self.after_cancel(self._search_timer)
            self._search_timer = self.after(250, self._perform_search)

        def _perform_search(self) -> None:
            self._search_timer = None
            self.refresh()

        def _move(self, direction: int) -> None:
            self.offset = max(0, self.offset + direction * PAGE_SIZE)
            self.refresh()

        def refresh(self) -> None:
            if self._search_timer is not None:
                self.after_cancel(self._search_timer)
                self._search_timer = None
            current = self.selected()
            self.table.delete(*self.table.get_children())
            page = self.app.repo.page(self.kind, self.offset, PAGE_SIZE, self.search.get())
            if self.offset and self.offset >= page["total"]:
                self.offset = max(0, (page["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                page = self.app.repo.page(self.kind, self.offset, PAGE_SIZE, self.search.get())
            for index, row in enumerate(page["rows"]):
                self.table.insert("", "end", iid=row["id"], values=tuple(row[field] for field in self.visible),
                                  tags=("stripe",) if index % 2 else ())
            self.page_label.set(f"{self.offset + 1 if page['total'] else 0}–{self.offset + len(page['rows'])} de {page['total']}")
            if current and self.table.exists(current):
                self.table.selection_set(current)

        def open_detail(self) -> None:
            key = self.selected()
            row = self.app.repo.get(self.kind, key) if key else None
            if not row:
                return
            dialog = tk.Toplevel(self)
            dialog.title(f"{LABELS.get(self.kind, self.kind)} — {key}")
            dialog.transient(self.winfo_toplevel())
            dialog.geometry("760x560")
            dialog.minsize(560, 420)
            dialog.columnconfigure(0, weight=1)
            dialog.rowconfigure(0, weight=1)
            notebook = ttk.Notebook(dialog)
            notebook.grid(row=0, column=0, sticky="nsew", padx=8, pady=8)

            info_tab = ttk.Frame(notebook, padding=10)
            notebook.add(info_tab, text="Información")
            info = tk.Text(info_tab, wrap="word", state="disabled",
                           font=("Segoe UI", 10), background=self.app.colors["surface"],
                           foreground=self.app.colors["ink"],
                           selectbackground=self.app.colors["text_select"],
                           padx=10, pady=8, borderwidth=0,
                           highlightthickness=1, highlightbackground=self.app.colors["line"])
            info_y = ttk.Scrollbar(info_tab, orient="vertical", command=info.yview)
            info.configure(yscrollcommand=info_y.set)
            info.pack(side="left", fill="both", expand=True)
            info_y.pack(side="right", fill="y")
            text = "\n".join(f"{LABELS.get(field, field)}: {row[field]}" for field in ENTITY_FIELDS[self.kind])
            if self.kind == "grupos":
                members = self.app.repo.related_page("membresias", key, "left", limit=0, active_only=True)["total"]
                products = self.app.repo.related_page("grupos_productos", key, "left", limit=0, active_only=True)["total"]
                text += f"\n\nIntegrantes: {members} | Productos vinculados: {products}"
            elif self.kind == "investigadores":
                groups = self.app.repo.related_page("membresias", key, "right", limit=0, active_only=True)["total"]
                products = self.app.repo.related_page("autorias", key, "right", limit=0, active_only=True)["total"]
                text += f"\n\nGrupos: {groups} | Productos de autoría: {products}"
            elif self.kind == "productos":
                authors = self.app.repo.related_page("autorias", key, "left", limit=0, active_only=True)["total"]
                groups = self.app.repo.related_page("grupos_productos", key, "right", limit=0, active_only=True)["total"]
                text += f"\n\nAutores: {authors} | Grupos: {groups}"
            info.configure(state="normal")
            info.insert("1.0", text)
            info.configure(state="disabled")

            products_tab = ttk.Frame(notebook, padding=10)
            notebook.add(products_tab, text="Productos")
            products_frame = ttk.Frame(products_tab)
            products_frame.pack(fill="both", expand=True)
            products_frame.columnconfigure(0, weight=1)
            products_frame.rowconfigure(0, weight=1)
            columns = ("id", "titulo", "anio", "tipologia", "categoria", "validacion")
            products_table = ttk.Treeview(products_frame, columns=columns, show="headings", selectmode="browse")
            products_table.tag_configure("stripe", background=self.app.colors["stripe"])
            for field in columns:
                products_table.heading(field, text=LABELS[field])
                products_table.column(field, width=150 if field in ("id", "anio") else 280, stretch=True)
            products_scroll = ttk.Scrollbar(products_frame, orient="vertical", command=products_table.yview)
            products_table.configure(yscrollcommand=products_scroll.set)
            products_table.grid(row=0, column=0, sticky="nsew")
            products_scroll.grid(row=0, column=1, sticky="ns")
            products = self._related_products(key)
            for index, product in enumerate(products):
                products_table.insert("", "end", iid=product["id"],
                                      values=tuple(product[field] for field in columns),
                                      tags=("stripe",) if index % 2 else ())
            if not products:
                ttk.Label(products_tab, text="Sin productos asociados").pack(pady=20)

            ttk.Button(dialog, text="Cerrar", command=dialog.destroy).grid(row=1, column=0, pady=(0, 10))
            dialog.bind("<Escape>", lambda _event: dialog.destroy())
            dialog.grab_set()

        def _related_products(self, key: str) -> list[dict[str, str]]:
            if self.kind == "grupos":
                links = self.app.repo.related("grupos_productos", key, "left")
            elif self.kind == "investigadores":
                links = self.app.repo.related("autorias", key, "right")
            elif self.kind == "productos":
                product = self.app.repo.get("productos", key)
                return [product] if product else []
            else:
                links = []
            products = []
            for link in links:
                if link.get("activo") != "1":
                    continue
                product = self.app.repo.get("productos", link["producto_id"])
                if product is not None:
                    products.append(product)
            return products

        def create(self) -> None:
            choices = {"grupo_id": [r["id"] for r in self.app.repo.page("grupos")["rows"]]}
            dialog = RecordDialog(self, self.kind, choices=choices,
                                  suggested_id=self.app.repo.suggest_id(self.kind))
            if dialog.result:
                self.app.mutate(lambda: self.app.repo.create(self.kind, dialog.result))

        def edit(self) -> None:
            key = self.selected()
            if not key:
                return
            dialog = RecordDialog(self, self.kind, self.app.repo.get(self.kind, key))
            if dialog.result:
                self.app.mutate(lambda: self.app.repo.update(self.kind, key, dialog.result))

        def open_source(self) -> None:
            key = self.selected()
            if key:
                row = self.app.repo.get(self.kind, key)
                if row and row["url"]:
                    self.app.import_url(row["url"])
                else:
                    messagebox.showinfo("Sin fuente", "Este registro no tiene URL de origen", parent=self)

        def toggle(self) -> None:
            key = self.selected()
            if key:
                self.app.mutate(lambda: self.app.repo.toggle(self.kind, key))

        def delete(self) -> None:
            key = self.selected()
            if key and messagebox.askyesno("Eliminar", f"¿Eliminar {key} definitivamente?", parent=self):
                self.app.mutate(lambda: self.app.repo.delete(self.kind, key))

        def restyle(self) -> None:
            self.table.tag_configure("stripe", background=self.app.colors["stripe"])

    class RelationTab(ttk.Frame):
        def __init__(self, parent: tk.Misc, app: "App", kind: str):
            super().__init__(parent, padding=8)
            self.app, self.kind = app, kind
            self.offset = 0
            actions = ttk.Frame(self)
            actions.pack(fill="x")
            for label, action in (("Vincular", self.create), ("Editar", self.edit),
                                  ("Activar/desactivar", self.toggle), ("Desvincular", self.delete)):
                ttk.Button(actions, text=label, command=action).pack(side="left", padx=3)
            self.fields = RELATION_FIELDS[kind]
            table_frame = ttk.Frame(self)
            table_frame.pack(fill="both", expand=True, pady=7)
            table_frame.columnconfigure(0, weight=1)
            table_frame.rowconfigure(0, weight=1)
            self.table = ttk.Treeview(table_frame, columns=self.fields, show="headings", selectmode="browse")
            self.table.tag_configure("stripe", background=self.app.colors["stripe"])
            for field in self.fields:
                self.table.heading(field, text=LABELS.get(field, field))
                self.table.column(field, width=150)
            table_y = ttk.Scrollbar(table_frame, orient="vertical", command=self.table.yview)
            table_x = ttk.Scrollbar(table_frame, orient="horizontal", command=self.table.xview)
            self.table.configure(yscrollcommand=table_y.set, xscrollcommand=table_x.set)
            self.table.grid(row=0, column=0, sticky="nsew")
            table_y.grid(row=0, column=1, sticky="ns")
            table_x.grid(row=1, column=0, sticky="ew")
            pager = ttk.Frame(self)
            pager.pack(fill="x")
            ttk.Button(pager, text="Anterior", command=lambda: self._move(-1)).pack(side="left")
            self.page_label = tk.StringVar()
            ttk.Label(pager, textvariable=self.page_label).pack(side="left", padx=10)
            ttk.Button(pager, text="Siguiente", command=lambda: self._move(1)).pack(side="left")

        def _move(self, direction: int) -> None:
            self.offset = max(0, self.offset + direction * PAGE_SIZE)
            self.refresh()

        def selected(self) -> tuple[str, str] | None:
            selection = self.table.selection()
            if not selection:
                return None
            return tuple(json.loads(selection[0]))  # type: ignore[return-value]

        def refresh(self) -> None:
            current = self.table.selection()
            self.table.delete(*self.table.get_children())
            left, right = RELATION_ENDS[self.kind][:2]
            page = self.app.repo.page(self.kind, self.offset, PAGE_SIZE)
            if self.offset and self.offset >= page["total"]:
                self.offset = max(0, (page["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                page = self.app.repo.page(self.kind, self.offset, PAGE_SIZE)
            for index, row in enumerate(page["rows"]):
                pair = json.dumps((row[left], row[right]))
                self.table.insert("", "end", iid=pair, values=tuple(row[field] for field in self.fields),
                                  tags=("stripe",) if index % 2 else ())
            self.page_label.set(f"{self.offset + 1 if page['total'] else 0}–{self.offset + len(page['rows'])} de {page['total']}")
            if current and self.table.exists(current[0]):
                self.table.selection_set(current[0])

        def create(self) -> None:
            left, right, left_kind, right_kind = RELATION_ENDS[self.kind]
            choices = {left: [r["id"] for r in self.app.repo.page(left_kind)["rows"]],
                       right: [r["id"] for r in self.app.repo.page(right_kind)["rows"]]}
            dialog = RecordDialog(self, self.kind, choices=choices)
            if dialog.result:
                self.app.mutate(lambda: self.app.repo.create(self.kind, dialog.result))

        def edit(self) -> None:
            key = self.selected()
            if not key:
                return
            dialog = RecordDialog(self, self.kind, self.app.repo.get(self.kind, key))
            if dialog.result:
                self.app.mutate(lambda: self.app.repo.update(self.kind, key, dialog.result))

        def toggle(self) -> None:
            key = self.selected()
            if key:
                self.app.mutate(lambda: self.app.repo.toggle(self.kind, key))

        def delete(self) -> None:
            key = self.selected()
            if key and messagebox.askyesno("Desvincular", "¿Eliminar esta relación?", parent=self):
                self.app.mutate(lambda: self.app.repo.delete(self.kind, key))

        def restyle(self) -> None:
            self.table.tag_configure("stripe", background=self.app.colors["stripe"])

    LIGHT_THEME = {
        "page": "#eef3f8", "surface": "#ffffff", "navy": "#13304d",
        "ink": "#23384c", "muted": "#587086", "line": "#d4e0ea",
        "teal": "#087f79", "hover": "#e1edf3", "strong": "#13304d",
        "header_fg": "#ffffff", "status_fg": "#d7e8ef",
        "accent_bg": "#087f79", "accent_fg": "#ffffff",
        "accent_active": "#096b67", "accent_pressed": "#075a57",
        "tab_bg": "#dce7ef", "heading_bg": "#e0eaf1", "heading_active": "#d1e2eb",
        "stripe": "#f3f8fc", "select_bg": "#13304d", "text_select": "#dce7ef",
        "chart_bg": "#ffffff", "chart_border": "#d4e0ea", "chart_title": "#18324f",
        "chart_empty": "#66788a", "chart_axis": "#93a5b5",
    }

    DARK_THEME = {
        "page": "#13171d", "surface": "#1b2129", "navy": "#1d2c3f",
        "ink": "#d8e0e8", "muted": "#8996a4", "line": "#2c3540",
        "teal": "#46c7b8", "hover": "#242d37", "strong": "#edf3f8",
        "header_fg": "#eef4f9", "status_fg": "#a9bccb",
        "accent_bg": "#2c9c90", "accent_fg": "#ffffff",
        "accent_active": "#36b1a3", "accent_pressed": "#247f74",
        "tab_bg": "#1e252e", "heading_bg": "#242c36", "heading_active": "#303a46",
        "stripe": "#171d25", "select_bg": "#2f5b7c", "text_select": "#35434f",
        "chart_bg": "#1b2129", "chart_border": "#2c3540", "chart_title": "#d8e0e8",
        "chart_empty": "#8996a4", "chart_axis": "#3f4c59",
    }

    class App:
        def __init__(self, root: tk.Tk, initial_dir: Path | None,
                     repository: Repository | CppRepository):
            self.root = root
            self.repo = repository
            self.backend_name = "C++" if isinstance(repository, CppRepository) else "Python"
            self.theme = "light"
            self.data_dir: Path | None = None
            self.initial_dir = initial_dir
            self._url_busy = False
            self.root.title("PEA-i UPC — Investigación")
            self.root.geometry("1280x850")
            self.root.minsize(900, 620)
            self._apply_theme("light")
            self._menu()
            self.sidebar = ttk.Frame(root, style="Sidebar.TFrame")
            self.sidebar.pack(side="left", fill="y")
            self.menu_expanded = True
            self.burger = ttk.Button(self.sidebar, text="☰", style="Burger.TButton",
                                     command=self._toggle_menu)
            self.burger.pack(anchor="n", padx=8, pady=(10, 4))
            self.sidebar_items = ttk.Frame(self.sidebar, style="Sidebar.TFrame")
            self.sidebar_items.pack(fill="both", expand=True)
            self.content = ttk.Frame(root)
            self.content.pack(side="left", fill="both", expand=True)
            self.content.rowconfigure(0, weight=1)
            self.content.columnconfigure(0, weight=1)

            self.pages: dict[str, tk.Widget] = {}
            self.dashboard = ttk.Frame(self.content, padding=10)
            self._build_dashboard()
            self.graphs = ttk.Frame(self.content, padding=10)
            self._build_graphs()
            self.entity_tabs: dict[str, EntityTab] = {}
            for kind in ("grupos", "investigadores", "productos", "planes"):
                self.entity_tabs[kind] = EntityTab(self.content, self, kind)
            relations_host = ttk.Frame(self.content, padding=10)
            relations = ttk.Notebook(relations_host)
            relations.pack(fill="both", expand=True)
            self.relation_tabs: dict[str, RelationTab] = {}
            for kind, title in (("membresias", "Integrantes"), ("autorias", "Autorías"),
                                ("grupos_productos", "Grupos / productos")):
                tab = RelationTab(relations, self, kind)
                self.relation_tabs[kind] = tab
                relations.add(tab, text=title)
            self.queue_tab = ttk.Frame(self.content, padding=10)
            self._build_queue()

            for page_id, title, page in (
                    ("dashboard", "Dashboard", self.dashboard),
                    ("graphs", "Gráficos", self.graphs),
                    ("grupos", "Grupos", self.entity_tabs["grupos"]),
                    ("investigadores", "Investigadores", self.entity_tabs["investigadores"]),
                    ("productos", "Productos", self.entity_tabs["productos"]),
                    ("planes", "Planes", self.entity_tabs["planes"]),
                    ("relaciones", "Relaciones", relations_host),
                    ("cola", "Cola de revisión", self.queue_tab)):
                self.pages[page_id] = page
                ttk.Button(self.sidebar_items, text=title, style="Nav.TButton",
                           command=lambda p=page_id: self._show_page(p)).pack(fill="x", padx=6, pady=2)
            self.theme_button = ttk.Button(self.sidebar_items, text="Modo oscuro", command=self._toggle_theme)
            self.theme_button.pack(side="bottom", fill="x", padx=6, pady=(10, 6))
            for page in self.pages.values():
                page.grid(row=0, column=0, sticky="nsew")

            self.status = tk.StringVar(value="Iniciando PEA-i...")
            self.refresh()
            self._show_page("dashboard")
            self.root.protocol("WM_DELETE_WINDOW", self.close)
            self.root.after(80, self.start_choice)

        def _show_page(self, page_id: str) -> None:
            self.pages[page_id].tkraise()
            if page_id == "graphs":
                self._redraw_charts()

        def _toggle_menu(self) -> None:
            if self.menu_expanded:
                self.sidebar_items.pack_forget()
            else:
                self.sidebar_items.pack(fill="both", expand=True)
            self.menu_expanded = not self.menu_expanded

        def _show_status(self) -> None:
            messagebox.showinfo("Estado de datos", self.status.get(), parent=self.root)

        def _apply_theme(self, theme: str) -> None:
            style = ttk.Style()
            if "clam" in style.theme_names():
                style.theme_use("clam")
            colors = LIGHT_THEME if theme == "light" else DARK_THEME
            self.theme = theme
            self.colors = colors
            self.root.configure(background=colors["page"])
            style.configure(".", font=("Segoe UI", 10), background=colors["page"], foreground=colors["ink"])
            style.configure("TFrame", background=colors["page"])
            style.configure("TLabel", background=colors["page"], foreground=colors["ink"])
            style.configure("Sidebar.TFrame", background=colors["navy"])
            style.configure("Nav.TButton", background=colors["navy"], foreground=colors["header_fg"],
                            bordercolor=colors["navy"], relief="flat", anchor="w",
                            padding=(14, 9), font=("Segoe UI", 10))
            style.map("Nav.TButton", background=[("active", colors["hover"]),
                                                 ("pressed", colors["hover"])],
                      foreground=[("active", colors["ink"] if theme == "light" else colors["header_fg"])])
            style.configure("Burger.TButton", background=colors["navy"], foreground=colors["header_fg"],
                            bordercolor=colors["navy"], relief="flat", padding=(12, 6),
                            font=("Segoe UI", 16))
            style.map("Burger.TButton", background=[("active", colors["hover"])],
                      foreground=[("active", colors["header_fg"])])
            style.configure("Heading.TLabel", font=("Segoe UI", 19, "bold"), foreground=colors["strong"])
            style.configure("Metric.TLabel", font=("Segoe UI", 17, "bold"), foreground=colors["teal"])
            style.configure("TButton", padding=(11, 7), background=colors["surface"],
                            foreground=colors["strong"], bordercolor=colors["line"], relief="flat")
            style.map("TButton", background=[("active", colors["hover"]), ("pressed", colors["line"])])
            style.configure("Accent.TButton", background=colors["accent_bg"], foreground=colors["accent_fg"],
                            bordercolor=colors["accent_bg"], padding=(12, 8))
            style.map("Accent.TButton", background=[("active", colors["accent_active"]),
                                                    ("pressed", colors["accent_pressed"])],
                      foreground=[("active", colors["accent_fg"])])
            style.configure("TNotebook", background=colors["page"], borderwidth=0, tabmargins=(10, 8, 0, 0))
            style.configure("TNotebook.Tab", background=colors["tab_bg"], foreground=colors["strong"],
                            padding=(15, 9), borderwidth=0)
            style.map("TNotebook.Tab", background=[("selected", colors["surface"]),
                                                    ("active", colors["hover"])],
                      foreground=[("selected", colors["teal"])])
            style.configure("TLabelframe", background=colors["page"], bordercolor=colors["line"],
                            relief="solid")
            style.configure("TLabelframe.Label", background=colors["page"], foreground=colors["strong"],
                            font=("Segoe UI", 10, "bold"))
            style.configure("TEntry", fieldbackground=colors["surface"], foreground=colors["ink"],
                            bordercolor=colors["line"], padding=5)
            style.configure("TCombobox", fieldbackground=colors["surface"], foreground=colors["ink"],
                            bordercolor=colors["line"], padding=5)
            style.map("TCombobox", fieldbackground=[("readonly", colors["surface"])],
                      foreground=[("readonly", colors["ink"])])
            style.configure("Treeview", background=colors["surface"], fieldbackground=colors["surface"],
                            foreground=colors["ink"], bordercolor=colors["line"], rowheight=30)
            style.map("Treeview", background=[("selected", colors["select_bg"])],
                      foreground=[("selected", colors["header_fg"])])
            style.configure("Treeview.Heading", background=colors["heading_bg"], foreground=colors["strong"],
                            bordercolor=colors["line"], relief="flat", padding=(7, 8),
                            font=("Segoe UI", 10, "bold"))
            style.map("Treeview.Heading", background=[("active", colors["heading_active"])])
            for menu in getattr(self, "_menus", ()):
                menu.configure(background=colors["surface"], foreground=colors["ink"],
                               activebackground=colors["hover"], activeforeground=colors["teal"],
                               borderwidth=0)
            self._restyle_widgets()

        def _toggle_theme(self) -> None:
            self._apply_theme("dark" if self.theme == "light" else "light")
            self.theme_button.configure(text="Modo claro" if self.theme == "dark" else "Modo oscuro")

        def _restyle_widgets(self) -> None:
            colors = self.colors
            for name in ("year_canvas", "type_canvas", "category_canvas", "validation_canvas"):
                canvas = getattr(self, name, None)
                if canvas is not None:
                    canvas.configure(bg=colors["chart_bg"], highlightbackground=colors["chart_border"])
            for name in ("stats_table",):
                table = getattr(self, name, None)
                if table is not None:
                    table.tag_configure("stripe", background=colors["stripe"])
            for tab in getattr(self, "entity_tabs", {}).values():
                tab.restyle()
            for tab in getattr(self, "relation_tabs", {}).values():
                tab.restyle()
            if hasattr(self, "queue_table"):
                self.queue_table.tag_configure("stripe", background=colors["stripe"])
            if hasattr(self, "current_stats"):
                self._redraw_charts()

        def _menu(self) -> None:
            menu = tk.Menu(self.root)
            files = tk.Menu(menu, tearoff=False)
            for label, command in (("Iniciar vacío", self.new_workspace),
                                   ("Abrir carpeta de datos...", self.open_folder),
                                   ("Cargar datos reales", self.load_real),
                                   ("Cargar demostración", self.load_demo),
                                   ("Guardar", self.save), ("Guardar como...", self.save_as),
                                   ("Exportar para hoja de cálculo...", self.export_spreadsheet),
                                   ("Salir", self.close)):
                files.add_command(label=label, command=command)
            menu.add_cascade(label="Archivo", menu=files)
            imports = tk.Menu(menu, tearoff=False)
            for label, command in (("Hoja de cálculo (CSV/Excel)...", self.import_csv),
                                   ("URL pública...", self.import_url),
                                   ("PDF de texto GrupLAC/CvLAC...", self.import_pdf),
                                   ("Documento Word...", self.import_docx)):
                imports.add_command(label=label, command=command)
            menu.add_cascade(label="Importar", menu=imports)
            edit = tk.Menu(menu, tearoff=False)
            edit.add_command(label="Deshacer", command=self.undo, accelerator="Ctrl+Z")
            edit.add_command(label="Vaciar historial", command=self.clear_history)
            menu.add_cascade(label="Editar", menu=edit)
            help_menu = tk.Menu(menu, tearoff=False)
            help_menu.add_command(label="Estado de datos", command=self._show_status)
            help_menu.add_command(label="Acerca de PEA-i", command=lambda: messagebox.showinfo(
                "Acerca de PEA-i", f"Programa Estadístico de Análisis de Investigación\n"
                f"Universidad Popular del Cesar\nMotor de datos: {self.backend_name}", parent=self.root))
            menu.add_cascade(label="Ayuda", menu=help_menu)
            self._menus = (menu, files, imports, edit, help_menu)
            self.root.bind_all("<Control-z>", lambda _event: self.undo())
            self.root.config(menu=menu)

        def _build_dashboard(self) -> None:
            ttk.Label(self.dashboard, text="Panorama de investigación", style="Heading.TLabel").pack(anchor="w")
            filters = ttk.LabelFrame(self.dashboard, text="Filtros", padding=8)
            filters.pack(fill="x", pady=8)
            self.view = tk.StringVar(value="Todos")
            self.scope_id = tk.StringVar()
            self._scope_kind: str | None = None
            self.period = tk.StringVar(value="Todos")
            self.start_year = tk.StringVar()
            self.end_year = tk.StringVar()
            self.category = tk.StringVar(value="Todas")
            self.validation = tk.StringVar(value="Todos")
            ttk.Label(filters, text="Vista").grid(row=0, column=0, padx=4, pady=3)
            view_box = ttk.Combobox(filters, textvariable=self.view, state="readonly", width=11,
                                    values=("Todos", "Grupo", "Investigador", "Producto"))
            view_box.grid(row=0, column=1, padx=4)
            view_box.bind("<<ComboboxSelected>>", lambda _event: self._scope_changed())
            ttk.Label(filters, text="ID").grid(row=0, column=2, padx=4)
            self.scope_box = ttk.Combobox(filters, textvariable=self.scope_id, state="readonly", width=15)
            self.scope_box.grid(row=0, column=3, padx=4)
            self.scope_box.bind("<<ComboboxSelected>>", lambda _event: self.refresh_dashboard())
            ttk.Label(filters, text="Ventana").grid(row=0, column=4, padx=4)
            period_box = ttk.Combobox(filters, textvariable=self.period, state="readonly", width=14,
                                      values=("Todos", "Últimos 2 años", "Últimos 5 años", "Personalizado"))
            period_box.grid(row=0, column=5, padx=4)
            period_box.bind("<<ComboboxSelected>>", lambda _event: self.refresh_dashboard())
            ttk.Label(filters, text="Desde").grid(row=1, column=0, padx=4)
            ttk.Entry(filters, textvariable=self.start_year, width=10).grid(row=1, column=1, sticky="w", padx=4)
            ttk.Label(filters, text="Hasta").grid(row=1, column=2, padx=4)
            ttk.Entry(filters, textvariable=self.end_year, width=10).grid(row=1, column=3, sticky="w", padx=4)
            ttk.Label(filters, text="Categoría").grid(row=1, column=4, padx=4)
            self.category_box = ttk.Combobox(filters, textvariable=self.category, state="readonly", width=14)
            self.category_box.grid(row=1, column=5, padx=4)
            self.category_box.bind("<<ComboboxSelected>>", lambda _event: self.refresh_dashboard())
            ttk.Label(filters, text="Validación").grid(row=2, column=0, padx=4, pady=3)
            status_box = ttk.Combobox(filters, textvariable=self.validation, state="readonly", width=13,
                                      values=("Todos", "pendiente", "validado", "rechazado"))
            status_box.grid(row=2, column=1, padx=4)
            status_box.bind("<<ComboboxSelected>>", lambda _event: self.refresh_dashboard())
            self.metric = tk.StringVar()
            ttk.Label(self.dashboard, textvariable=self.metric, style="Metric.TLabel").pack(anchor="w", pady=5)
            self.note = tk.StringVar(value="Cada producto se cuenta una vez dentro de la vista y el filtro.")
            ttk.Label(self.dashboard, textvariable=self.note).pack(anchor="w")
            ttk.Label(self.dashboard, text="Productos de la vista").pack(anchor="w")
            columns = ("id", "titulo", "anio", "tipologia", "categoria", "validacion")
            table_frame = ttk.Frame(self.dashboard)
            table_frame.pack(fill="both", expand=True)
            self.stats_table = ttk.Treeview(table_frame, columns=columns, show="headings", height=6)
            self.stats_table.tag_configure("stripe", background=self.colors["stripe"])
            column_widths = {"id": 155, "titulo": 430, "anio": 80, "tipologia": 230,
                             "categoria": 130, "validacion": 125}
            for field in columns:
                self.stats_table.heading(field, text=LABELS[field])
                self.stats_table.column(field, width=column_widths[field], stretch=False)
            x_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.stats_table.xview)
            y_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.stats_table.yview)
            self.stats_table.configure(xscrollcommand=x_scroll.set, yscrollcommand=y_scroll.set)
            self.stats_table.grid(row=0, column=0, sticky="nsew")
            y_scroll.grid(row=0, column=1, sticky="ns")
            x_scroll.grid(row=1, column=0, sticky="ew")
            table_frame.rowconfigure(0, weight=1)
            table_frame.columnconfigure(0, weight=1)
            self.stats_table.bind("<Double-1>", self._open_product)
            stats_pager = ttk.Frame(self.dashboard)
            stats_pager.pack(fill="x")
            ttk.Button(stats_pager, text="Anterior", command=lambda: self._stats_move(-1)).pack(side="left")
            self.stats_page_label = tk.StringVar()
            ttk.Label(stats_pager, textvariable=self.stats_page_label).pack(side="left", padx=10)
            ttk.Button(stats_pager, text="Siguiente", command=lambda: self._stats_move(1)).pack(side="left")
            self.stats_offset = 0
            self._stats_filter: tuple[Any, ...] | None = None
            self.current_stats: dict[str, Any] = {"por_anio": {}, "por_tipologia": {},
                                                  "por_categoria": {}, "por_validacion": {}}

        def _build_graphs(self) -> None:
            ttk.Label(self.graphs, text="Visualización de productos", style="Heading.TLabel").pack(anchor="w")
            charts = ttk.Notebook(self.graphs)
            charts.pack(fill="both", expand=True, pady=8)
            self.year_canvas = tk.Canvas(charts, height=200, bg=self.colors["chart_bg"],
                                         highlightthickness=1, highlightbackground=self.colors["chart_border"])
            self.type_canvas = tk.Canvas(charts, height=240, bg=self.colors["chart_bg"],
                                         highlightthickness=1, highlightbackground=self.colors["chart_border"])
            self.category_canvas = tk.Canvas(charts, height=240, bg=self.colors["chart_bg"],
                                             highlightthickness=1, highlightbackground=self.colors["chart_border"])
            self.validation_canvas = tk.Canvas(charts, height=240, bg=self.colors["chart_bg"],
                                               highlightthickness=1, highlightbackground=self.colors["chart_border"])
            charts.add(self.year_canvas, text="Por año")
            charts.add(self.type_canvas, text="Por tipología")
            charts.add(self.category_canvas, text="Por categoría")
            charts.add(self.validation_canvas, text="Por validación")
            for canvas in (self.year_canvas, self.type_canvas, self.category_canvas, self.validation_canvas):
                canvas.bind("<Configure>", lambda _event: self._redraw_charts())

        def _build_queue(self) -> None:
            controls = ttk.Frame(self.queue_tab)
            controls.pack(fill="x")
            for label, command in (("Encolar producto", self.enqueue), ("Procesar frente", self.process_queue),
                                   ("Descartar frente", self.discard_queue)):
                ttk.Button(controls, text=label, command=command).pack(side="left", padx=4)
            self.queue_count = tk.StringVar()
            ttk.Label(controls, textvariable=self.queue_count).pack(side="right")
            columns = ("id", "producto_id", "motivo", "creado")
            table_frame = ttk.Frame(self.queue_tab)
            table_frame.pack(fill="both", expand=True, pady=8)
            table_frame.columnconfigure(0, weight=1)
            table_frame.rowconfigure(0, weight=1)
            self.queue_table = ttk.Treeview(table_frame, columns=columns, show="headings")
            self.queue_table.tag_configure("stripe", background=self.colors["stripe"])
            for field in columns:
                self.queue_table.heading(field, text=LABELS.get(field, field))
                self.queue_table.column(field, width=160)
            table_y = ttk.Scrollbar(table_frame, orient="vertical", command=self.queue_table.yview)
            table_x = ttk.Scrollbar(table_frame, orient="horizontal", command=self.queue_table.xview)
            self.queue_table.configure(yscrollcommand=table_y.set, xscrollcommand=table_x.set)
            self.queue_table.grid(row=0, column=0, sticky="nsew")
            table_y.grid(row=0, column=1, sticky="ns")
            table_x.grid(row=1, column=0, sticky="ew")
            queue_pager = ttk.Frame(self.queue_tab)
            queue_pager.pack(fill="x")
            ttk.Button(queue_pager, text="Anterior", command=lambda: self._queue_move(-1)).pack(side="left")
            self.queue_page_label = tk.StringVar()
            ttk.Label(queue_pager, textvariable=self.queue_page_label).pack(side="left", padx=10)
            ttk.Button(queue_pager, text="Siguiente", command=lambda: self._queue_move(1)).pack(side="left")
            self.queue_offset = 0

        def _queue_move(self, direction: int) -> None:
            self.queue_offset = max(0, self.queue_offset + direction * PAGE_SIZE)
            self.refresh_queue()

        def refresh_queue(self) -> None:
            page = self.repo.queue_page(self.queue_offset, PAGE_SIZE)
            if self.queue_offset and self.queue_offset >= page["total"]:
                self.queue_offset = max(0, (page["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                page = self.repo.queue_page(self.queue_offset, PAGE_SIZE)
            self.queue_table.delete(*self.queue_table.get_children())
            for index, job in enumerate(page["rows"]):
                self.queue_table.insert("", "end", iid=job["id"],
                                        values=tuple(job[name] for name in ("id", "producto_id", "motivo", "creado")),
                                        tags=("stripe",) if index % 2 else ())
            self.queue_count.set(f"Pendientes: {page['total']}")
            self.queue_page_label.set(f"{self.queue_offset + 1 if page['total'] else 0}–{self.queue_offset + len(page['rows'])} de {page['total']}")

        def enqueue(self) -> None:
            ids = [p["id"] for p in self.repo.page("productos", limit=20)["rows"]]
            if not ids:
                messagebox.showinfo("Sin productos", "Cree un producto primero", parent=self.root)
                return
            product_id = simpledialog.askstring("Encolar revisión", "ID de producto:\n" + ", ".join(ids[:20]), parent=self.root)
            if not product_id:
                return
            reason = simpledialog.askstring("Motivo", "Motivo de la revisión:", parent=self.root)
            if reason is None:
                return
            self.mutate(lambda: self.repo.enqueue_review(product_id.strip(), reason))

        def process_queue(self) -> None:
            job = self.repo.queue_front()
            if not job:
                messagebox.showinfo("Cola", "No hay revisiones pendientes", parent=self.root)
                return
            status = simpledialog.askstring("Procesar frente", f"Producto {job['producto_id']}\nEstado: validado, rechazado o pendiente", parent=self.root)
            if status is None:
                return
            observation = simpledialog.askstring("Observación", "Razón o evidencia del resultado:", parent=self.root)
            if observation is None:
                return
            self.mutate(lambda: self.repo.process_review(status.strip().lower(), observation))

        def discard_queue(self) -> None:
            job = self.repo.queue_front()
            if job and messagebox.askyesno("Descartar frente", f"¿Descartar revisión de {job['producto_id']}?", parent=self.root):
                self.mutate(self.repo.discard_review)

        def _stats_move(self, direction: int) -> None:
            self.stats_offset = max(0, self.stats_offset + direction * PAGE_SIZE)
            self.refresh_dashboard()

        def _redraw_charts(self) -> None:
            self._draw_bars(self.year_canvas, chart_values(self.current_stats.get("por_anio", {})),
                            "Productos por año", "#2585a2")
            self._draw_donut(self.type_canvas, chart_values(self.current_stats.get("por_tipologia", {})),
                             "Productos por tipología")
            self._draw_donut(self.category_canvas, chart_values(self.current_stats.get("por_categoria", {})),
                             "Productos por categoría")
            self._draw_donut(self.validation_canvas, chart_values(self.current_stats.get("por_validacion", {})),
                             "Productos por validación")

        def _draw_bars(self, canvas: tk.Canvas, values: dict[str, int], title: str, color: str) -> None:
            canvas.delete("all")
            width = max(canvas.winfo_width(), 300)
            height = max(canvas.winfo_height(), 190)
            canvas.create_text(12, 13, anchor="w", text=title, font=("TkDefaultFont", 10, "bold"),
                               fill=self.colors["chart_title"])
            if not values:
                canvas.create_text(width / 2, height / 2, text="Sin valores registrados",
                                   fill=self.colors["chart_empty"])
                return
            if title == "Productos por año":
                items = sorted(values.items(), key=lambda pair: pair[0])
                if len(items) > 12:
                    items = [("Antes", sum(value for _, value in items[:-11]))] + items[-11:]
            else:
                items = sorted(values.items(), key=lambda pair: (-pair[1], pair[0]))
                if len(items) > 12:
                    items = items[:11] + [("Otros", sum(value for _, value in items[11:]))]
            maximum = max(value for _, value in items)
            left, right, top, bottom = 38, width - 15, 35, height - 45
            bar_space = (right - left) / len(items)
            for index, (name, value) in enumerate(items):
                x0 = left + index * bar_space + 4
                x1 = left + (index + 1) * bar_space - 4
                bar_height = (bottom - top) * value / maximum
                canvas.create_rectangle(x0, bottom - bar_height, x1, bottom, fill=color, outline="")
                canvas.create_text((x0 + x1) / 2, bottom - bar_height - 10, text=str(value),
                                   fill=self.colors["chart_title"])
                short_name = name if len(name) <= 11 else name[:10] + "…"
                canvas.create_text((x0 + x1) / 2, bottom + 12, text=short_name, width=bar_space - 4,
                                   font=("TkDefaultFont", 8), fill=self.colors["chart_title"])
            canvas.create_line(left, bottom, right, bottom, fill=self.colors["chart_axis"])

        def _draw_donut(self, canvas: tk.Canvas, values: dict[str, int], title: str) -> None:
            canvas.delete("all")
            width = max(canvas.winfo_width(), 360)
            height = max(canvas.winfo_height(), 240)
            canvas.create_text(12, 13, anchor="w", text=title, font=("TkDefaultFont", 10, "bold"),
                               fill=self.colors["chart_title"])
            if not values:
                canvas.create_text(width / 2, height / 2, text="Sin valores registrados",
                                   fill=self.colors["chart_empty"])
                return
            items = sorted(values.items(), key=lambda pair: -pair[1])
            if len(items) > 8:
                items = items[:7] + [("Otros", sum(value for _, value in items[7:]))]
            palette = ("#2585a2", "#087f79", "#d49544", "#b4556d", "#7a68b0",
                       "#5b8f5a", "#c77d3c", "#8a9096")
            total = sum(value for _, value in items)
            cx = width * 0.28
            cy = height / 2
            radius = min(height * 0.40, width * 0.20)
            thickness = radius * 0.46
            start = 90.0
            for index, (_, value) in enumerate(items):
                extent = -360.0 * value / total
                canvas.create_arc(cx - radius, cy - radius, cx + radius, cy + radius,
                                  start=start, extent=extent, style="arc",
                                  outline=palette[index % len(palette)], width=thickness)
                start += extent
            canvas.create_text(cx, cy - 10, text=str(total), font=("TkDefaultFont", 17, "bold"),
                               fill=self.colors["chart_title"])
            canvas.create_text(cx, cy + 11, text="productos", font=("TkDefaultFont", 9),
                               fill=self.colors["chart_empty"])
            legend_x = width * 0.52
            legend_y = max(28, height / 2 - (len(items) - 1) * 13)
            for index, (label, value) in enumerate(items):
                y = legend_y + index * 27
                pct = round(100 * value / total)
                canvas.create_rectangle(legend_x, y - 6, legend_x + 13, y + 6,
                                        fill=palette[index % len(palette)], outline="")
                short = label if len(label) <= 24 else label[:23] + "…"
                canvas.create_text(legend_x + 20, y, anchor="w", text=f"{short} · {pct}% ({value})",
                                   font=("TkDefaultFont", 9), fill=self.colors["chart_title"])

        def _scope_changed(self) -> None:
            kind = {"Grupo": "grupos", "Investigador": "investigadores", "Producto": "productos"}.get(self.view.get())
            choices = [row["id"] for row in self.repo.page(kind)["rows"]] if kind else []
            self.scope_box.configure(values=choices, state="normal" if kind else "disabled")
            if kind != self._scope_kind or (kind and not self.scope_id.get()):
                self.scope_id.set(choices[0] if kind and choices else "")
            self._scope_kind = kind
            self.refresh_dashboard()

        def refresh_dashboard(self) -> None:
            try:
                start = end = None
                period = self.period.get()
                if period in ("Últimos 2 años", "Últimos 5 años"):
                    years = 2 if period == "Últimos 2 años" else 5
                    start, end = date.today().year - years + 1, date.today().year
                elif period == "Personalizado":
                    start = int(self.start_year.get()) if self.start_year.get().strip() else None
                    end = int(self.end_year.get()) if self.end_year.get().strip() else None
                    if start is not None and end is not None and start > end:
                        raise DataError("El año inicial supera al final")
                    if any(value is not None and not 1900 <= value <= date.today().year + 1 for value in (start, end)):
                        raise DataError("Año fuera de rango")
                category = "" if self.category.get() == "Todas" else self.category.get()
                status = "" if self.validation.get() == "Todos" else self.validation.get()
                filters = (self.view.get(), self.scope_id.get(), start, end, category, status)
                if filters != self._stats_filter:
                    self.stats_offset = 0
                    self._stats_filter = filters
                result = self.repo.statistics(*filters, offset=self.stats_offset, limit=PAGE_SIZE)
                if self.stats_offset and self.stats_offset >= result["total"]:
                    self.stats_offset = max(0, (result["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                    result = self.repo.statistics(*filters, offset=self.stats_offset, limit=PAGE_SIZE)
            except (ValueError, DataError) as exc:
                self.status.set(f"Filtro inválido: {exc}")
                return
            self.current_stats = result
            active_groups = self.overview["active_groups"]
            active_people = self.overview["active_people"]
            self.metric.set(f"{result['total']} productos únicos  •  {active_groups} grupos habilitados  •  {active_people} perfiles habilitados")
            if result["total"]:
                missing = [f"{label}: {result[field].get('Sin dato', 0)} sin dato"
                           for label, field in (("Año", "por_anio"), ("Tipología", "por_tipologia"),
                                                ("Categoría", "por_categoria"))
                           if result[field].get("Sin dato", 0)]
                states = ", ".join(f"{key} {value}" for key, value in result["por_validacion"].items())
                self.note.set("Estados: " + states + ("   •   " + "   •   ".join(missing) if missing else ""))
            else:
                self.note.set("Sin productos para estos filtros")
            self.stats_table.delete(*self.stats_table.get_children())
            for index, row in enumerate(result["productos"]):
                self.stats_table.insert("", "end", iid=row["id"],
                                        values=tuple(row[field] for field in ("id", "titulo", "anio", "tipologia", "categoria", "validacion")),
                                        tags=("stripe",) if index % 2 else ())
            self.stats_page_label.set(f"{self.stats_offset + 1 if result['total'] else 0}–{self.stats_offset + len(result['productos'])} de {result['total']}")
            self._redraw_charts()

        def _open_product(self, _event: tk.Event) -> None:
            selection = self.stats_table.selection()
            if selection:
                self._show_page("productos")
                self.entity_tabs["productos"].search.set(selection[0])
                self.entity_tabs["productos"].refresh()
                self.entity_tabs["productos"].table.selection_set(selection[0])
                self.entity_tabs["productos"].table.see(selection[0])
                self.entity_tabs["productos"].open_detail()

        def refresh(self) -> None:
            for tab in self.entity_tabs.values():
                tab.refresh()
            for tab in self.relation_tabs.values():
                tab.refresh()
            self.refresh_queue()
            self.overview = self.repo.summary()
            categories = ["Todas"] + sorted(name for name in self.overview["categories"] if name != "Sin dato")
            self.category_box.configure(values=categories)
            if self.category.get() not in categories:
                self.category.set("Todas")
            self._scope_changed()
            self.root.title("PEA-i UPC — Investigación" + (" *" if self.repo.dirty else ""))

        def mutate(self, action: Callable[[], Any]) -> Any:
            try:
                result = action()
            except (DataError, OSError, ValueError) as exc:
                messagebox.showerror("No se pudo completar", str(exc), parent=self.root)
                return None
            self.refresh()
            self.status.set("Cambio en memoria. Guarde para conservarlo al cerrar.")
            return result

        def _may_discard(self) -> bool:
            if not self.repo.dirty:
                return True
            choice = messagebox.askyesnocancel("Cambios sin guardar", "¿Guardar cambios antes de continuar?", parent=self.root)
            if choice is None:
                return False
            if choice:
                return self.save()
            return True

        def start_choice(self) -> None:
            if self.initial_dir and (self.initial_dir / "manifest.csv").exists():
                self._load(self.initial_dir)
                return
            self.load_real()

        def new_workspace(self) -> None:
            if self._may_discard():
                if isinstance(self.repo, CppRepository):
                    self.repo.reset()
                else:
                    self.repo = Repository()
                self.data_dir = None
                self.refresh()
                self.status.set("Espacio vacío. Guarde en una carpeta antes de salir.")

        def _load(self, directory: Path, demo: bool = False) -> None:
            recovered = False
            try:
                if isinstance(self.repo, CppRepository):
                    self.repo.load(directory)
                else:
                    new_repo = load_repository(directory)
            except (DataError, OSError, ValueError, UnicodeError, json.JSONDecodeError) as exc:
                backup = directory / ".backup"
                if backup.is_dir() and messagebox.askyesno("Datos dañados", f"{exc}\n\n¿Abrir la copia anterior de .backup?", parent=self.root):
                    try:
                        if isinstance(self.repo, CppRepository):
                            self.repo.load(backup)
                        else:
                            new_repo = load_repository(backup)
                    except (DataError, OSError, ValueError, UnicodeError, json.JSONDecodeError) as backup_exc:
                        messagebox.showerror("Copia inválida", str(backup_exc), parent=self.root)
                        return
                    directory = backup
                    recovered = True
                else:
                    messagebox.showerror("No se pudo cargar", str(exc), parent=self.root)
                    return
            if not isinstance(self.repo, CppRepository):
                self.repo = new_repo
            data_root = Path(__file__).resolve().parents[2] / "data"
            demo = demo or directory.resolve() in ((data_root / "demo").resolve(), (data_root / "real").resolve())
            self.data_dir = None if demo or recovered else directory
            self.refresh()
            self.status.set(f"Datos cargados de {directory}" + (" (guardar como copia)" if demo or recovered else ""))

        def open_folder(self) -> None:
            if not self._may_discard():
                return
            folder = filedialog.askdirectory(parent=self.root, title="Carpeta con manifest.csv")
            if folder:
                self._load(Path(folder))

        def load_demo(self) -> None:
            if self._may_discard():
                self._load(Path(__file__).resolve().parents[2] / "data" / "demo", demo=True)

        def load_real(self) -> None:
            if self._may_discard():
                self._load(Path(__file__).resolve().parents[2] / "data" / "real", demo=True)

        def save(self) -> bool:
            if self.data_dir is None:
                return self.save_as()
            try:
                if isinstance(self.repo, CppRepository):
                    self.repo.save(self.data_dir)
                else:
                    save_repository(self.repo, self.data_dir)
            except (OSError, DataError) as exc:
                messagebox.showerror("No se pudo guardar", str(exc), parent=self.root)
                return False
            self.refresh()
            self.status.set(f"Guardado en {self.data_dir}")
            return True

        def save_as(self) -> bool:
            folder = filedialog.askdirectory(parent=self.root, title="Carpeta de trabajo para guardar")
            if not folder:
                return False
            chosen = Path(folder)
            data_root = Path(__file__).resolve().parents[2] / "data"
            if chosen.resolve() in ((data_root / "demo").resolve(), (data_root / "real").resolve()):
                messagebox.showerror("Muestra protegida", "Guarde una copia en otra carpeta", parent=self.root)
                return False
            if (chosen / "manifest.csv").exists() and chosen != self.data_dir:
                if not messagebox.askyesno("Reemplazar datos", "La carpeta ya contiene datos PEA-i. ¿Reemplazarlos?", parent=self.root):
                    return False
            previous = self.data_dir
            self.data_dir = chosen
            if not self.save():
                self.data_dir = previous
                return False
            return True

        def export_spreadsheet(self) -> None:
            folder = filedialog.askdirectory(parent=self.root, title="Carpeta para exportar a hoja de cálculo")
            if not folder:
                return
            try:
                export_repository_spreadsheet(self.repo, Path(folder))
            except (OSError, DataError) as exc:
                messagebox.showerror("No se pudo exportar", str(exc), parent=self.root)
                return
            self.status.set(f"Exportado para hoja de cálculo en {folder}")

        def close(self) -> None:
            if self._may_discard():
                self.root.destroy()

        def undo(self) -> None:
            try:
                undone = self.repo.undo()
            except DataError as exc:
                messagebox.showerror("No se pudo deshacer", str(exc), parent=self.root)
                return
            if undone:
                self.refresh()
                self.status.set("Última acción deshecha")
            else:
                self.status.set("No hay acciones para deshacer")

        def clear_history(self) -> None:
            if self.repo.history_size() and messagebox.askyesno("Vaciar historial", "¿Eliminar las acciones para deshacer?", parent=self.root):
                self.mutate(self.repo.clear_history)

        def import_csv(self) -> None:
            path = filedialog.askopenfilename(parent=self.root, title="Importar hoja de cálculo",
                                              filetypes=[("Hojas de cálculo", "*.csv *.xlsx"),
                                                         ("CSV", "*.csv"), ("Excel", "*.xlsx")])
            if not path:
                return
            path = Path(path)
            if path.suffix.lower() == ".xlsx":
                try:
                    with tempfile.TemporaryDirectory() as temporary:
                        self.import_csv_path(_xlsx_to_csv(path, Path(temporary) / "hoja.csv"))
                except (DataError, OSError) as exc:
                    messagebox.showerror("Excel inválido", str(exc), parent=self.root)
                return
            self.import_csv_path(path)

        def import_csv_path(self, path: Path) -> None:
            kinds = ", ".join(ALL_FIELDS)
            kind = simpledialog.askstring("Tipo de datos", f"Nombre del tipo CSV:\n{kinds}",
                                          initialvalue=path.stem, parent=self.root)
            if not kind:
                return
            kind = kind.strip()
            if kind not in ALL_FIELDS:
                messagebox.showerror("Tipo desconocido", kinds, parent=self.root)
                return
            try:
                if isinstance(self.repo, CppRepository):
                    total, expected, errors = self.repo.preview_csv(path, kind)
                else:
                    total, expected, errors = preview_csv_merge(self.repo, path, kind)
            except (DataError, OSError, UnicodeError) as exc:
                messagebox.showerror("CSV inválido", str(exc), parent=self.root)
                return
            summary = f"{total} filas de {kind}\nAceptables: {expected}\nRechazadas: {total - expected}"
            if errors:
                summary += "\n\n" + "\n".join(errors[:8])
            if not messagebox.askyesno("Vista previa", summary + "\n\n¿Importar las aceptables?", parent=self.root):
                return
            try:
                if isinstance(self.repo, CppRepository):
                    accepted, errors = self.repo.import_csv(path, kind)
                else:
                    accepted, errors = merge_csv(self.repo, path, kind)
            except (DataError, OSError, UnicodeError) as exc:
                messagebox.showerror("No se pudo importar", str(exc), parent=self.root)
                return
            self.refresh()
            messagebox.showinfo("Resultado de importación", f"Aceptadas: {accepted}\nRechazadas: {total - accepted}\n\n" + "\n".join(errors[:20]), parent=self.root)

        def _accept_external(self, kind: str, row: dict[str, str]) -> None:
            existing = self.repo.get(kind, row["id"]) if row.get("id") else None
            initial = {**existing, **{key: value for key, value in row.items() if value}} if existing else row
            if existing and existing.get("fuente") and row.get("fuente") and row["fuente"] not in existing["fuente"]:
                initial["fuente"] = existing["fuente"] + "; " + row["fuente"]
            dialog = RecordDialog(self.root, kind, initial, external=existing is None)
            if dialog.result:
                if existing:
                    self.mutate(lambda: self.repo.update(kind, existing["id"], dialog.result))
                else:
                    self.mutate(lambda: self.repo.create(kind, dialog.result))

        def import_url(self, initial_url: str | None = None) -> None:
            if self._url_busy:
                self.status.set("Ya hay una consulta URL en curso")
                return
            url = initial_url or simpledialog.askstring("Consultar URL", "Pegue una URL pública HTTP o HTTPS:", parent=self.root)
            if not url:
                return
            self.status.set("Descargando página pública...")
            self._url_busy = True
            result_queue: queue.Queue[WebPreview | Exception] = queue.Queue(maxsize=1)

            def fetch() -> None:
                try:
                    result_queue.put(fetch_web_page(url.strip()))
                except Exception as exc:
                    result_queue.put(exc)

            def check_result() -> None:
                try:
                    alive = self.root.winfo_exists()
                except tk.TclError:
                    return
                if not alive:
                    return
                try:
                    result = result_queue.get_nowait()
                except queue.Empty:
                    self.root.after(100, check_result)
                    return
                self._url_busy = False
                if isinstance(result, Exception):
                    messagebox.showerror("Consulta URL", str(result), parent=self.root)
                    self.status.set("La consulta no modificó los datos")
                    return
                self.status.set(f"Página consultada: {result.title}")
                self._show_web_preview(result)

            threading.Thread(target=fetch, daemon=True).start()
            self.root.after(100, check_result)

        def _show_web_preview(self, preview: WebPreview, source: str | None = None) -> None:
            dialog = tk.Toplevel(self.root)
            dialog.title("Vista previa de URL pública")
            dialog.transient(self.root)
            dialog.geometry("850x650")
            dialog.minsize(650, 450)
            dialog.columnconfigure(0, weight=1)
            dialog.rowconfigure(4, weight=1)
            ttk.Label(dialog, text=preview.title, style="Heading.TLabel", wraplength=780,
                      padding=(12, 10)).grid(row=0, column=0, sticky="w")
            ttk.Label(dialog, text=f"Formato: {preview.format}  •  Fuente: {source or _web_source(preview.url)}",
                      padding=(12, 0)).grid(row=1, column=0, sticky="w")
            url_box = ttk.Entry(dialog)
            url_box.insert(0, preview.url)
            url_box.configure(state="readonly")
            url_box.grid(row=2, column=0, sticky="ew", padx=12, pady=6)
            note = "Datos detectados para revisar: " + (preview.suggested_kind or "ningún registro de investigación")
            if preview.related_members or preview.related_products:
                note += (f"; {len(preview.related_members)} integrantes, {len(preview.related_products)} productos"
                         f" y {len(preview.related_authorships)} autorías coincidentes con el censo")
                if preview.related_plan:
                    note += "; plan estratégico"
            note += ". Los gráficos usan productos guardados y solo campos con valores."
            ttk.Label(dialog, text=note, wraplength=780, padding=(12, 4)).grid(row=3, column=0, sticky="w")
            body = ttk.Frame(dialog, padding=(12, 4))
            body.grid(row=4, column=0, sticky="nsew")
            body.columnconfigure(0, weight=1)
            body.rowconfigure(0, weight=1)
            shown = tk.Text(body, wrap="word", state="normal")
            scroll = ttk.Scrollbar(body, orient="vertical", command=shown.yview)
            shown.configure(yscrollcommand=scroll.set)
            shown.grid(row=0, column=0, sticky="nsew")
            scroll.grid(row=0, column=1, sticky="ns")
            details = ["METADATOS DECLARADOS POR LA PÁGINA"]
            details.extend(f"{key}: {value}" for key, value in preview.metadata.items())
            if not preview.metadata:
                details.append("Ninguno")
            details.extend(("", "TEXTO VISIBLE (primeras 300 líneas)", *preview.lines))
            if preview.tables:
                details.append("\nTABLAS DETECTADAS (hasta 5 tablas y 100 filas por tabla)")
                for index, table in enumerate(preview.tables, start=1):
                    details.append(f"Tabla {index}")
                    details.extend(" | ".join(row) for row in table)
                    details.append("")
            shown.insert("1.0", "\n".join(details))
            shown.configure(state="disabled")
            actions = ttk.Frame(dialog, padding=10)
            actions.grid(row=5, column=0, sticky="ew")

            def create(kind: str) -> None:
                row = (preview.suggested_row or {}).copy() if preview.suggested_kind == kind else {}
                row.setdefault("id", new_id({"grupos": "G", "investigadores": "I", "productos": "P"}[kind]))
                if source is None:
                    row["url"] = _redact_url(preview.url)
                    row.setdefault("fuente", _web_source(preview.url))
                else:
                    row.setdefault("fuente", source)
                dialog.destroy()
                self._accept_external(kind, row)

            def import_group_profile() -> None:
                dialog.destroy()
                proposed = (preview.suggested_row or {}).copy()
                if not proposed:
                    return
                reviewed = None
                if self.repo.get("grupos", proposed["id"]) is None:
                    review = RecordDialog(self.root, "grupos", proposed, external=True)
                    if not review.result:
                        return
                    reviewed = review.result
                self.status.set("Incorporando ficha GrupLAC...")
                self.root.update_idletasks()
                try:
                    counts = import_gruplac_preview(self.repo, preview, reviewed)
                except (DataError, OSError, ValueError) as exc:
                    messagebox.showerror("Ficha GrupLAC", str(exc), parent=self.root)
                    return
                self.refresh()
                self.status.set("Ficha incorporada en memoria. Guarde una copia para conservarla."
                                if self.data_dir is None else "Ficha incorporada en memoria. Use Guardar para conservarla.")
                summary = "\n".join(f"{label}: {counts[key]}" for key, label in (
                    ("grupos", "Grupos"), ("investigadores", "Investigadores"),
                    ("membresias", "Membresías"), ("planes", "Planes"),
                    ("productos", "Productos"), ("grupos_productos", "Vínculos grupo-producto"),
                    ("autorias", "Autorías"), ("existentes", "Ya existentes")))
                if counts["errors"]:
                    summary += f"\nErrores: {len(counts['errors'])}\n" + "\n".join(counts["errors"][:5])
                messagebox.showinfo("Ficha GrupLAC", summary, parent=self.root)

            for kind, label in (("grupos", "Crear grupo"), ("investigadores", "Crear investigador"),
                                ("productos", "Crear producto")):
                ttk.Button(actions, text=label, command=lambda name=kind: create(name)).pack(side="left", padx=3)
            if preview.related_members or preview.related_products or preview.related_plan:
                ttk.Button(actions, text="Importar datos detectados", style="Accent.TButton",
                           command=import_group_profile).pack(side="left", padx=3)
            if preview.csv_bytes is not None:
                def import_downloaded_csv() -> None:
                    dialog.destroy()
                    with tempfile.TemporaryDirectory() as temporary:
                        path = Path(temporary) / "descarga.csv"
                        path.write_bytes(preview.csv_bytes or b"")
                        self.import_csv_path(path)
                ttk.Button(actions, text="Importar CSV PEA-i", command=import_downloaded_csv).pack(side="left", padx=3)
            ttk.Button(actions, text="Cerrar", command=dialog.destroy).pack(side="right", padx=3)
            dialog.bind("<Escape>", lambda _event: dialog.destroy())
            dialog.grab_set()

        def import_pdf(self) -> None:
            path = filedialog.askopenfilename(parent=self.root, title="PDF de GrupLAC o CvLAC", filetypes=[("PDF", "*.pdf")])
            if not path:
                return
            url = simpledialog.askstring("Origen del PDF", "URL original pública de la ficha GrupLAC/CvLAC:", parent=self.root)
            if not url:
                return
            try:
                kind, row = parse_scienti_pdf(Path(path), url.strip())
                row["fuente"] = f"PDF {Path(path).name}, consultado {date.today().isoformat()}"
            except (DataError, OSError, ValueError) as exc:
                messagebox.showerror("Importación PDF", str(exc), parent=self.root)
                return
            self._accept_external(kind, row)

        def import_docx(self) -> None:
            path = filedialog.askopenfilename(parent=self.root, title="Documento Word",
                                              filetypes=[("Word", "*.docx")])
            if not path:
                return
            path = Path(path)
            try:
                html = extract_docx_html(path)
            except DataError as exc:
                messagebox.showerror("Documento Word", str(exc), parent=self.root)
                return
            preview = parse_web_page(html, path.resolve().as_uri(), "HTML")
            self._show_web_preview(preview, source=f"DOCX {path.name}, consultado {date.today().isoformat()}")

    root = tk.Tk()
    repository: Repository | CppRepository | None = None
    try:
        repository, fallback_reason = open_gui_repository(
            python_backend=python_backend, cpp_binary=cpp_binary)
        app = App(root, data_dir, repository)
        if fallback_reason:
            app.status.set(f"Motor Python activo: {fallback_reason}")
        root.mainloop()
    except DataError as exc:
        messagebox.showerror("Backend C++", str(exc), parent=root)
        raise
    finally:
        if isinstance(repository, CppRepository):
            repository.close()
        try:
            root.destroy()
        except tk.TclError:
            pass


def main() -> None:
    parser = argparse.ArgumentParser(description="PEA-i UPC con interfaz Tkinter")
    parser.add_argument("--data-dir", type=Path, help="Carpeta PEA-i que se abre al iniciar")
    parser.add_argument("--cpp-binary", type=Path, help="Ejecutable C++ compilado (opcional)")
    parser.add_argument("--python-backend", action="store_true",
                        help="Usar la implementación Python independiente en vez del backend C++")
    args = parser.parse_args()
    try:
        run_gui(args.data_dir, python_backend=args.python_backend, cpp_binary=args.cpp_binary)
    except DataError as exc:
        parser.exit(1, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
