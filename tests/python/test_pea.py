"""Pruebas del núcleo, ejecutables con python -m unittest discover -s tests/python."""

import csv
from email.message import Message
import hashlib
import importlib.util
import os
from pathlib import Path
import shutil
import tempfile
import time
import unittest
from unittest.mock import patch


SOURCE = Path(__file__).resolve().parents[2] / "src" / "python" / "Taller2_REMR.py"
SPEC = importlib.util.spec_from_file_location("pea", SOURCE)
pea = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(pea)


def sample() -> pea.Repository:
    repo = pea.Repository()
    repo.create("grupos", {"id": "G1", "nombre": "Grupo uno"})
    repo.create("grupos", {"id": "G2", "nombre": "Grupo dos"})
    repo.create("investigadores", {"id": "I1", "nombre": "Ana"})
    repo.create("investigadores", {"id": "I2", "nombre": "Beto"})
    repo.create("productos", {"id": "P1", "titulo": "Artículo, con coma", "anio": "2025", "tipologia": "Artículo", "categoria": "A1"})
    repo.create("productos", {"id": "P2", "titulo": "Trabajo\nmultilínea", "anio": "2024", "tipologia": "Libro", "categoria": "B"})
    repo.create("membresias", {"grupo_id": "G1", "investigador_id": "I1", "rol": "Líder"})
    repo.create("membresias", {"grupo_id": "G1", "investigador_id": "I2", "rol": "Integrante"})
    repo.create("membresias", {"grupo_id": "G2", "investigador_id": "I1", "rol": "Integrante"})
    repo.create("autorias", {"producto_id": "P1", "investigador_id": "I1", "orden": "1"})
    repo.create("autorias", {"producto_id": "P1", "investigador_id": "I2", "orden": "2"})
    repo.create("grupos_productos", {"grupo_id": "G1", "producto_id": "P1"})
    repo.create("grupos_productos", {"grupo_id": "G2", "producto_id": "P1"})
    repo.create("grupos_productos", {"grupo_id": "G2", "producto_id": "P2"})
    repo.create("planes", {"id": "PL1", "grupo_id": "G1", "nombre": "Plan", "objetivo": "Publicar"})
    return repo


class StartupTest(unittest.TestCase):
    def test_windows_bundle_matches_cpp_source(self):
        root = SOURCE.parents[2]
        self.assertEqual(pea.bundled_windows_cpp(root), root / "bin/windows/pea_cpp.exe")
        with tempfile.TemporaryDirectory() as folder:
            checkout = Path(folder)
            for name in ("CMakeLists.txt", "src/cpp/Taller2_REMR.cpp"):
                copied = checkout / name
                copied.parent.mkdir(parents=True, exist_ok=True)
                content = (root / name).read_bytes().replace(b"\r\n", b"\n")
                copied.write_bytes(content.replace(b"\n", b"\r\n"))
            binary = checkout / "bin/windows/pea_cpp.exe"
            binary.parent.mkdir(parents=True)
            binary.write_bytes(b"MZ")
            binary.with_suffix(".source-sha256").write_bytes(
                (root / "bin/windows/pea_cpp.source-sha256").read_bytes())
            (checkout / "bin/windows/pea_cpp.exe.sha256").write_bytes(
                (hashlib.sha256(b"MZ").hexdigest() + "\n").encode("ascii"))
            self.assertEqual(pea.bundled_windows_cpp(checkout), binary)

    def test_python_fallback_and_explicit_cpp_error(self):
        with patch.object(pea, "CppRepository", side_effect=pea.DataError("CMake falló\nNMake no disponible")):
            repository, notice = pea.open_gui_repository()
            self.assertIsInstance(repository, pea.Repository)
            self.assertEqual(notice, "CMake falló")
            with self.assertRaisesRegex(pea.DataError, "CMake falló"):
                pea.open_gui_repository(cpp_binary=Path("backend.exe"))


class StructuresTest(unittest.TestCase):
    def test_doubly_linked_list_boundaries(self):
        linked = pea.DoublyLinkedList()
        self.assertEqual(list(linked), [])
        for key in "ABC":
            linked.append({"id": key})
        self.assertEqual([v["id"] for v in linked.backwards()], list("CBA"))
        linked.remove("A")
        linked.remove("C")
        linked.remove("B")
        self.assertEqual(linked.length, 0)
        self.assertIsNone(linked.head)
        self.assertIsNone(linked.tail)

    def test_multilist_two_directions_and_unlink(self):
        repo = sample()
        links = repo.relations["membresias"]
        self.assertEqual({r["investigador_id"] for r in links.by_left("G1")}, {"I1", "I2"})
        self.assertEqual({r["grupo_id"] for r in links.by_right("I1")}, {"G1", "G2"})
        repo.delete("membresias", ("G1", "I1"))
        self.assertEqual([r["investigador_id"] for r in links.by_left("G1")], ["I2"])
        self.assertEqual([r["grupo_id"] for r in links.by_right("I1")], ["G2"])

    def test_queue_fifo_stack_lifo(self):
        queue = pea.LinkedQueue()
        queue.enqueue({"id": "A"})
        queue.enqueue({"id": "B"})
        self.assertEqual(queue.dequeue()["id"], "A")
        self.assertEqual(queue.dequeue()["id"], "B")
        self.assertIsNone(queue.dequeue())
        stack = pea.LinkedStack()
        stack.push("A")
        stack.push("B")
        self.assertEqual(stack.pop(), "B")
        self.assertEqual(stack.pop(), "A")
        self.assertIsNone(stack.pop())


class DomainTest(unittest.TestCase):
    def test_counts_dedupe_year_and_views(self):
        repo = sample()
        self.assertEqual(repo.statistics()["total"], 2)
        self.assertEqual(repo.statistics(view="Grupo", selected_id="G2")["total"], 2)
        self.assertEqual(repo.statistics(view="Investigador", selected_id="I1")["total"], 1)
        self.assertEqual(repo.statistics(view="Grupo", selected_id="G1", start=2025, end=2025)["total"], 1)
        self.assertEqual(repo.statistics(view="Grupo", selected_id="")["total"], 0)
        self.assertEqual(repo.statistics(category="A1")["total"], 1)

    def test_integrity_validation_and_undo(self):
        repo = sample()
        with self.assertRaises(pea.DataError):
            repo.delete("grupos", "G1")
        with self.assertRaises(pea.DataError):
            repo.create("autorias", {"producto_id": "P1", "investigador_id": "missing"})
        with self.assertRaises(pea.DataError):
            repo.update("productos", "P1", {"categoria": "A2"})
        repo.update("productos", "P1", {"categoria": "A2", "observacion": "Revisado por comité"})
        self.assertEqual(repo.get("productos", "P1")["categoria"], "A2")
        self.assertTrue(repo.undo())
        self.assertEqual(repo.get("productos", "P1")["categoria"], "A1")
        repo.toggle("productos", "P1")
        self.assertEqual(repo.statistics()["total"], 1)

    def test_review_queue_and_undo(self):
        repo = sample()
        repo.enqueue_review("P1", "Verificar")
        repo.enqueue_review("P2", "Verificar")
        self.assertEqual(repo.queue.peek()["producto_id"], "P1")
        repo.process_review("validado", "Evidencia suficiente")
        self.assertEqual(repo.queue.peek()["producto_id"], "P2")
        self.assertEqual(repo.get("productos", "P1")["validacion"], "validado")
        repo.undo()
        self.assertEqual(repo.queue.peek()["producto_id"], "P1")
        self.assertEqual(repo.get("productos", "P1")["validacion"], "pendiente")

    def test_save_load_utf8_and_history(self):
        repo = sample()
        repo.enqueue_review("P1", "Tildes: acción")
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            pea.save_repository(repo, path)
            loaded = pea.load_repository(path)
            self.assertEqual(loaded.snapshot(), repo.snapshot())
            self.assertEqual(loaded.history.length, repo.history.length)
            self.assertTrue(loaded.undo())
            self.assertEqual(loaded.queue.length, 0)
            self.assertEqual(loaded.get("productos", "P2")["titulo"], "Trabajo\nmultilínea")

    def test_previous_save_can_be_recovered(self):
        repo = sample()
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            pea.save_repository(repo, path)
            repo.update("grupos", "G1", {"nombre": "Nuevo nombre"})
            pea.save_repository(repo, path)
            (path / "grupos.csv").write_text("archivo dañado", encoding="utf-8")
            with self.assertRaises(pea.DataError):
                pea.load_repository(path)
            recovered = pea.load_repository(path / ".backup")
            self.assertEqual(recovered.get("grupos", "G1")["nombre"], "Grupo uno")

    def test_product_year_derived_from_date(self):
        repo = pea.Repository()
        row = repo.create("productos", {"id": "P-DATE", "titulo": "Con fecha", "fecha": "2025-03-12"})
        self.assertEqual(row["anio"], "2025")
        with self.assertRaises(pea.DataError):
            repo.create("productos", {"id": "P-BAD", "titulo": "Fechas distintas", "anio": "2024", "fecha": "2025-03-12"})

    def test_csv_partial_merge_and_bad_headers(self):
        repo = sample()
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "productos.csv"
            with path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=pea.ENTITY_FIELDS["productos"])
                writer.writeheader()
                writer.writerow({"id": "P3", "titulo": "Nuevo", "anio": "2026"})
                writer.writerow({"id": "P1", "titulo": "Duplicado", "anio": "2026"})
                writer.writerow({"id": "P4", "titulo": "Año imposible", "anio": "1800"})
            total, anticipated, preview_errors = pea.preview_csv_merge(repo, path, "productos")
            self.assertEqual((total, anticipated, len(preview_errors)), (3, 1, 2))
            self.assertIsNone(repo.get("productos", "P3"))
            accepted, errors = pea.merge_csv(repo, path, "productos")
            self.assertEqual(accepted, 1)
            self.assertEqual(len(errors), 2)
            self.assertIsNotNone(repo.get("productos", "P3"))
            repo.undo()
            self.assertIsNone(repo.get("productos", "P3"))

    def test_scienti_parsing(self):
        group_url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=0001"
        html = "<html><body><h1>Grupo de prueba</h1><h2>Datos básicos</h2><table><tr><td>Líder</td><td>Ana Pérez</td></tr><tr><td>Clasificación</td><td>A1</td></tr></table></body></html>"
        kind, group = pea.parse_scienti_html(html, group_url)
        self.assertEqual((kind, group["nombre"], group["responsable"], group["categoria"]),
                         ("grupos", "Grupo de prueba", "Ana Pérez", "A1"))
        cv_url = "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0042"
        kind, person = pea.parse_scienti_html("<h1>Hoja de vida</h1><p>Categoría</p><p>Investigador Junior (IJ) con vigencia</p><p>Nombre</p><p>Juana Álvarez</p>", cv_url)
        self.assertEqual((kind, person["nombre"], person["categoria"]),
                         ("investigadores", "Juana Álvarez", "Investigador Junior (IJ)"))
        with self.assertRaises(pea.DataError):
            pea.parse_scienti_html(html, "https://example.org/gruplac/?nro=0001")

    def test_scienti_download_uses_page_charset(self):
        url = "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0042"
        html = ('<meta charset="iso-8859-1"><h1>Hoja de vida</h1><p>Categoría</p>'
                '<p>Investigador Junior (IJ) con vigencia</p><p>Nombre</p><p>Juana Álvarez</p>').encode("iso-8859-1")

        class Response:
            def __init__(self):
                self.headers = Message()
                self.headers["Content-Type"] = "text/html"

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def geturl(self):
                return url

            def read(self, _limit):
                return html

        class Opener:
            def __init__(self, response):
                self.response = response

            def open(self, _request, timeout):
                return self.response

        public_address = [(2, 1, 6, "", ("93.184.215.14", 443))]
        with patch.object(pea.socket, "getaddrinfo", return_value=public_address):
            with patch.object(pea, "build_opener", return_value=Opener(Response())):
                preview = pea.fetch_web_page(url)
        self.assertEqual(preview.suggested_kind, "investigadores")
        self.assertEqual(preview.suggested_row["nombre"], "Juana Álvarez")

    def test_other_web_pages_are_previewed_without_invented_records(self):
        url = "https://example.org/about"
        html = """<html><head><title>Observatorio público</title>
        <meta name="description" content="Investigación regional"></head>
        <body><h1>Datos disponibles</h1><p>Hay 12 informes.</p>
        <table><tr><th>Año</th><th>Informes</th></tr><tr><td>2024</td><td>12</td></tr></table>
        <script>document.write('999 productos');</script></body></html>"""
        preview = pea.parse_web_page(html, url)
        self.assertEqual(preview.title, "Observatorio público")
        self.assertIn("Hay 12 informes.", preview.lines)
        self.assertNotIn("999 productos", " ".join(preview.lines))
        self.assertEqual(preview.metadata["description"], "Investigación regional")
        self.assertEqual(preview.tables[0][1], ("2024", "12"))
        self.assertIsNone(preview.suggested_kind)
        self.assertIsNone(preview.suggested_row)
        self.assertEqual(pea.chart_values({"Sin dato": 3}), {})

    def test_declared_scholarly_article_only_populates_declared_fields(self):
        url = "https://journal.example.org/article/42"
        html = """<html><head><title>Revista de ciencia</title>
        <meta name="citation_title" content="Artículo de prueba">
        <meta name="citation_publication_date" content="2024-05-14">
        <meta name="citation_doi" content="10.1000/prueba"></head>
        <body><h1>Artículo de prueba</h1></body></html>"""
        preview = pea.parse_web_page(html, url)
        self.assertEqual(preview.suggested_kind, "productos")
        self.assertEqual(preview.suggested_row["titulo"], "Artículo de prueba")
        self.assertEqual(preview.suggested_row["anio"], "2024")
        self.assertEqual(preview.suggested_row["doi"], "10.1000/prueba")
        self.assertNotIn("categoria", preview.suggested_row)
        self.assertNotIn("tipologia", preview.suggested_row)
        self.assertEqual(preview.suggested_row["url"], url)
        schema_only = pea.parse_web_page('''<script type="application/ld+json">
        {"@type":"https://schema.org/ScholarlyArticle","headline":"Segundo artículo",
         "datePublished":"2023-07-01"}</script>''', url)
        self.assertEqual(schema_only.suggested_row["titulo"], "Segundo artículo")
        self.assertEqual(schema_only.suggested_row["anio"], "2023")
        self.assertEqual(schema_only.suggested_row["doi"], "")

    def test_public_url_fetch_formats_limits_and_private_redirects(self):
        class Response:
            def __init__(self, url, body, media_type):
                self.url, self.body = url, body
                self.headers = Message()
                self.headers["Content-Type"] = media_type

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def geturl(self):
                return self.url

            def read(self, limit):
                return self.body[:limit]

        class Opener:
            def __init__(self, response):
                self.response = response

            def open(self, _request, timeout):
                self_timeout = timeout
                self_test.assertEqual(self_timeout, 15)
                return self.response

        self_test = self
        public_address = [(2, 1, 6, "", ("93.184.215.14", 443))]
        with patch.object(pea.socket, "getaddrinfo", return_value=public_address):
            with patch.object(pea, "build_opener", return_value=Opener(Response(
                "https://example.org/data.csv", b'id,titulo,anio\nP-1,"Uno, dos",2024\n', "text/csv; charset=utf-8"))):
                preview = pea.fetch_web_page("https://example.org/data.csv")
            self.assertEqual(preview.format, "CSV")
            self.assertEqual(preview.metadata["Columnas CSV"], "id, titulo, anio")
            self.assertIn("Uno, dos", preview.lines[1])
            self.assertIsNotNone(preview.csv_bytes)
            with patch.object(pea, "build_opener", return_value=Opener(Response(
                "http://127.0.0.1/private", b"<html>secret</html>", "text/html"))):
                with self.assertRaises(pea.DataError):
                    pea.fetch_web_page("https://example.org")
            with patch.object(pea, "build_opener", return_value=Opener(Response(
                "https://example.org/huge", b"x" * 8_000_001, "text/plain"))):
                with self.assertRaises(pea.DataError):
                    pea.fetch_web_page("https://example.org/huge")
        with self.assertRaises(pea.DataError):
            pea.fetch_web_page("http://127.0.0.1/private")
        with self.assertRaises(pea.DataError):
            pea.PublicRedirects().redirect_request(None, None, 302, "Found", {}, "http://127.0.0.1/private")

    def test_gruplac_articles_require_declared_title_and_year(self):
        url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
        html = """<h1>Grupo de prueba</h1><h2>Datos básicos</h2>
        <h2>Artículos publicados</h2>
        <p>1.- Publicado en revista especializada: Estudio verificable</p>
        <p>Colombia, Revista X ISSN: 1234-5678, 2024 vol:2, DOI:10.1000/uno</p>
        <p>Autores: Persona A</p>
        <p>2.- Publicado en revista especializada: Estudio verificable</p>
        <p>Colombia, Revista X ISSN: 1234-5678, 2024 vol:2, DOI:10.1000/uno</p>
        <p>3.- Publicado en revista especializada: Sin fecha</p>
        <p>Revista Y ISSN: 9999-9999, DOI:</p>
        <h2>Libros publicados</h2>
        <p>1.- Libro: No es artículo</p>"""
        rows = pea.parse_gruplac_articles(html, url)
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["titulo"], rows[0]["anio"], rows[0]["doi"]),
                         ("Estudio verificable", "2024", "10.1000/uno"))
        self.assertEqual(rows[0]["categoria"], "")
        self.assertEqual(rows[0]["validacion"], "pendiente")
        self.assertEqual(pea.parse_web_page(html, url).related_products[0]["id"], rows[0]["id"])

    def test_gruplac_roster_plan_products_and_repeat_import(self):
        url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
        profile = "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0042"
        html = f"""<h1>Grupo de prueba</h1><h2>Datos básicos</h2>
        <h2>Plan Estratégico</h2><p>Plan de trabajo: publicar libros y software.</p>
        <h2>Estado del arte</h2><p>Investigación previa.</p>
        <h2>Objetivos</h2><p>Compartir conocimiento.</p><h2>Retos</h2>
        <h2>Integrantes del grupo</h2><table>
        <tr><th>Nombre</th><th>Vinculación</th><th>Horas</th><th>Inicio - Fin</th></tr>
        <tr><td><a href="{profile}">1.- Ana Pérez</a></td><td>Investigador</td><td>20</td><td>2020/1 - Actual</td></tr>
        <tr><td>2.- Luis Gómez</td><td>Investigador</td><td>10</td><td>2018/1 - 2021/12</td></tr>
        </table>
        <h2>Artículos publicados</h2>
        <p>1.- Publicado en revista especializada: Estudio verificable</p>
        <p>Colombia, Revista X, 2024 vol:2, DOI:10.1000/uno</p>
        <p>Autores: Ana Perez, Persona Externa</p>
        <h2>Libros publicados</h2>
        <p>1.- Libro resultado de investigación: Libro probado</p>
        <p>Colombia, Editorial X, 2023.</p>
        <p>Autores: Luis Gómez</p>
        <h2>Capítulos de libro publicados</h2>"""
        preview = pea.parse_web_page(html, url)
        self.assertEqual(len(preview.related_members), 2)
        self.assertEqual(preview.related_members[0][0]["url"], profile)
        self.assertEqual([membership["activo"] for _, membership in preview.related_members], ["1", "0"])
        self.assertEqual(len(preview.related_products), 2)
        self.assertEqual({row["tipologia"] for row in preview.related_products},
                         {"Publicado en revista especializada", "Libro resultado de investigación"})
        self.assertIn("publicar libros y software", preview.related_plan["actividad"])
        self.assertEqual(len(preview.related_authorships), 2)
        for factory in (pea.Repository, lambda: pea.CppRepository(SOURCE.parents[2] / "build" / "pea_cpp")):
            if factory is not pea.Repository and not (SOURCE.parents[2] / "build" / "pea_cpp").exists():
                continue
            repo = factory()
            try:
                counts = pea.import_gruplac_preview(repo, preview)
                self.assertEqual([counts[kind] for kind in ("grupos", "investigadores", "membresias",
                                                           "planes", "productos", "grupos_productos", "autorias")],
                                 [1, 2, 2, 1, 2, 2, 2])
                self.assertEqual(counts["errors"], [])
                self.assertEqual(repo.statistics(view="Investigador", selected_id="I-0042")["total"], 1)
                again = pea.import_gruplac_preview(repo, preview)
                self.assertEqual(sum(again[kind] for kind in counts if kind not in ("existentes", "errors")), 0)
                self.assertEqual(again["errors"], [])
            finally:
                if isinstance(repo, pea.CppRepository):
                    repo.close()

    def test_real_dataset_loads_with_expected_relations(self):
        real = pea.load_repository(SOURCE.parents[2] / "data" / "real")
        self.assertEqual(real.statistics()["total"], 6363)
        self.assertEqual(real.statistics(view="Grupo", selected_id="G-00000000002099")["total"], 197)
        self.assertEqual(real.statistics(view="Investigador", selected_id="I-0000494917")["total"], 15)
        self.assertEqual(len(real.rows("grupos")), 66)
        self.assertEqual(len(real.rows("investigadores")), 2736)
        self.assertEqual(len(real.rows("membresias")), 3291)
        self.assertEqual(len(real.rows("planes")), 66)
        self.assertEqual(len(real.rows("autorias")), 9703)
        self.assertEqual(len(real.rows("grupos_productos")), 6735)
        self.assertTrue(all(not row["categoria"] for row in real.rows("productos")))
        self.assertEqual(sum(bool(row["categoria"]) for row in real.rows("investigadores")), 274)
        self.assertEqual(real.get("investigadores", "I-0000494917")["categoria"], "Investigador Asociado (I)")

    @unittest.skipUnless(shutil.which("pdftotext"), "Poppler no instalado")
    def test_text_pdf_import(self):
        # PDF pequeño generado directamente para probar el adaptador Poppler.
        content = b"BT /F1 12 Tf 50 750 Td (GrupLAC - Plataforma SCienTI - Colombia) Tj 0 -25 Td (Grupo PDF de prueba) Tj 0 -25 Td (Datos basicos) Tj ET"
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        ]
        data = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for index, obj in enumerate(objects, start=1):
            offsets.append(len(data))
            data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
        xref = len(data)
        data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
        for offset in offsets[1:]:
            data.extend(f"{offset:010d} 00000 n \n".encode())
        data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "profile.pdf"
            path.write_bytes(data)
            url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
            kind, group = pea.parse_scienti_pdf(path, url)
        self.assertEqual((kind, group["nombre"]), ("grupos", "Grupo PDF de prueba"))

    def test_demo_contract(self):
        demo = pea.load_repository(SOURCE.parents[2] / "data" / "demo")
        self.assertEqual(demo.statistics()["total"], 4)
        self.assertEqual(demo.statistics(start=2025, end=2026)["total"], 2)
        self.assertEqual(demo.statistics(start=2022, end=2026)["total"], 3)
        self.assertEqual(demo.statistics(view="Grupo", selected_id="G-DEMO-1")["total"], 3)
        self.assertEqual(demo.queue.peek()["producto_id"], "P-DEMO-2")


class SecurityTest(unittest.TestCase):
    def test_public_url_rejects_private_credentials_and_scheme(self):
        for bad in ("http://127.0.0.1/", "http://10.0.0.1/", "http://localhost/x",
                    "http://[::1]/", "ftp://example.org/x",
                    "http://user:pass@example.org/x"):
            with self.assertRaises(pea.DataError):
                pea._validate_public_url(bad)

    def test_public_url_rejects_nonstandard_port(self):
        with self.assertRaises(pea.DataError):
            pea._validate_public_url("https://example.org:8080/x")

    def test_public_url_rejects_private_resolution(self):
        mixed = [(2, 1, 6, "", ("93.184.215.14", 443)),
                 (2, 1, 6, "", ("127.0.0.1", 443))]
        with patch.object(pea.socket, "getaddrinfo", return_value=mixed):
            with self.assertRaises(pea.DataError):
                pea._validate_public_url("https://example.org/x")

    def test_resolve_public_addrs_dedupes(self):
        addresses = [(2, 1, 6, "", ("93.184.215.14", 443)),
                     (2, 1, 6, "", ("93.184.215.14", 443))]
        with patch.object(pea.socket, "getaddrinfo", return_value=addresses):
            self.assertEqual(pea._resolve_public_addrs("example.org", 443), ["93.184.215.14"])

    def test_public_opener_disables_proxy_and_pins(self):
        opener = pea._build_public_opener()
        self.assertTrue(any(isinstance(h, pea._PinnedHTTPHandler) for h in opener.handlers))
        self.assertTrue(any(isinstance(h, pea._PinnedHTTPSHandler) for h in opener.handlers))
        for handler in opener.handlers:
            self.assertFalse(getattr(handler, "proxies", {}))

    def test_redact_url_drops_sensitive_params(self):
        url = ("https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp"
               "?nro=0001&token=SECRET&auth=abc#frag")
        self.assertEqual(pea._redact_url(url),
                         "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=0001")
        self.assertEqual(pea._redact_url("https://example.org/a?token=x"), "https://example.org/a")

    def test_csv_neutralizes_formula_cells(self):
        self.assertEqual(pea._neutralize_cell("=SUM(A1)"), "'=SUM(A1)")
        self.assertEqual(pea._neutralize_cell("+cmd"), "'+cmd")
        self.assertEqual(pea._neutralize_cell("-2+3"), "'-2+3")
        self.assertEqual(pea._neutralize_cell("@cmd"), "'@cmd")
        self.assertEqual(pea._neutralize_cell("normal"), "normal")
        self.assertEqual(pea._neutralize_cell(""), "")

    def test_spreadsheet_export_neutralizes(self):
        repo = sample()
        repo.create("grupos", {"id": "G3", "nombre": '=HYPERLINK("http://x")'})
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            pea.export_repository_spreadsheet(repo, path)
            text = (path / "grupos.csv").read_text(encoding="utf-8")
            self.assertIn("'=HYPERLINK", text)

    def test_saved_csv_neutralizes_and_roundtrips_formula_text(self):
        repo = pea.Repository()
        repo.create("grupos", {"id": "G1", "nombre": "=1+1"}, remember=False)
        repo.create("productos", {"id": "P1", "titulo": "'=texto"}, remember=False)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            pea.save_repository(repo, path)
            with (path / "grupos.csv").open(encoding="utf-8", newline="") as stream:
                self.assertEqual(next(csv.DictReader(stream))["nombre"], "'=1+1")
            with (path / "productos.csv").open(encoding="utf-8", newline="") as stream:
                self.assertEqual(next(csv.DictReader(stream))["titulo"], "''=texto")
            loaded = pea.load_repository(path)
            self.assertEqual(loaded.get("grupos", "G1")["nombre"], "=1+1")
            self.assertEqual(loaded.get("productos", "P1")["titulo"], "'=texto")

    def test_legacy_v1_variants_load(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            pea.save_repository(sample(), path)
            with (path / "manifest.csv").open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(("version", "guardado"))
                writer.writerow(("1", "2026-09-24T00:00:00"))
            groups = pea._csv_read(path / "grupos.csv", pea.ENTITY_FIELDS["grupos"], encoded=True)
            with (path / "grupos.csv").open("w", encoding="utf-8", newline="") as stream:
                fields = ("id", "nombre", "sigla", *pea.ENTITY_FIELDS["grupos"][2:])
                writer = csv.DictWriter(stream, fieldnames=fields)
                writer.writeheader()
                writer.writerows(groups)
            self.assertEqual(pea.load_repository(path).get("grupos", "G1")["nombre"], "Grupo uno")
            products = pea._csv_read(path / "productos.csv", pea.ENTITY_FIELDS["productos"], encoded=True)
            with (path / "productos.csv").open("w", encoding="utf-8", newline="") as stream:
                fields = tuple(field for field in pea.ENTITY_FIELDS["productos"] if field != "validacion")
                writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(products)
            (path / "cola_validacion.csv").unlink()
            loaded = pea.load_repository(path)
            self.assertEqual(loaded.get("productos", "P1")["validacion"], "pendiente")
            self.assertEqual(loaded.queue_size(), 0)

    def test_extract_pdf_text_caps_output(self):
        class Stream:
            def __init__(self, data):
                self._data = data

            def read(self, size):
                chunk = self._data[:size]
                self._data = self._data[size:]
                return chunk

            def close(self):
                pass

        class Process:
            def __init__(self, stdout):
                self.stdout = Stream(stdout)
                self.stderr = Stream("")
                self.returncode = 0

            def kill(self):
                pass

            def wait(self, timeout=None):
                return 0

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "x.pdf"
            path.write_bytes(b"%PDF-1.4")
            with patch.object(pea.shutil, "which", return_value="/usr/bin/pdftotext"):
                with patch.object(pea.subprocess, "Popen",
                                  return_value=Process("x" * (pea.MAX_PDF_TEXT + 1))):
                    with self.assertRaises(pea.DataError):
                        pea.extract_pdf_text(path)

    @unittest.skipIf(os.name == "nt", "El ejecutable falso usa un shebang POSIX")
    def test_pdf_wall_timeout_stops_stalled_converter(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            converter = folder / "pdftotext"
            converter.write_text("#!/usr/bin/env python3\nimport time\ntime.sleep(3)\n", encoding="utf-8")
            converter.chmod(0o755)
            pdf = folder / "documento.pdf"
            pdf.write_bytes(b"%PDF-1.4")
            with patch.dict(os.environ, {"PATH": str(folder) + os.pathsep + os.environ["PATH"]}):
                with patch.object(pea, "PDF_TIMEOUT", 1):
                    started = time.monotonic()
                    with self.assertRaisesRegex(pea.DataError, "tiempo"):
                        pea.extract_pdf_text(pdf)
                    self.assertLess(time.monotonic() - started, 2.5)

    def test_field_length_limit(self):
        repo = pea.Repository()
        with self.assertRaises(pea.DataError):
            repo.create("grupos", {"id": "G1", "nombre": "x" * (pea.MAX_FIELD_LEN + 1)})

    def test_csv_row_limit(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "productos.csv"
            with path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=pea.ENTITY_FIELDS["productos"])
                writer.writeheader()
                for index in range(10):
                    writer.writerow({"id": f"P{index}", "titulo": "T"})
            with patch.object(pea, "MAX_CSV_ROWS", 5):
                with self.assertRaises(pea.DataError):
                    pea._csv_read(path, pea.ENTITY_FIELDS["productos"])

    def test_deeply_nested_jsonld_does_not_crash(self):
        nested = '{"@type":"ScholarlyArticle","@graph":' + '{"@graph":' * 5000 + '{}' + '}' * 5000 + '}'
        html = f'<script type="application/ld+json">{nested}</script>'
        preview = pea.parse_web_page(html, "https://example.org/x")
        self.assertIsNone(preview.suggested_kind)


def _minimal_docx(text: str) -> bytes:
    content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/word/document.xml" '
                     'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                     '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>')
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body><w:p><w:r><w:t>' + text + '</w:t></w:r></w:p></w:body></w:document>')
    import io
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
    return buffer.getvalue()


class ImportFormatsTest(unittest.TestCase):
    def test_xlsx_to_csv_roundtrip(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no instalado")
        workbook = openpyxl.Workbook()
        workbook.active.append(["id", "nombre"])
        workbook.active.append(["1", "Ana"])
        workbook.active.append(["2", "Beto"])
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "datos.xlsx"
            target = Path(temp) / "datos.csv"
            workbook.save(source)
            pea._xlsx_to_csv(source, target)
            rows = pea._csv_read(target, ("id", "nombre"))
            self.assertEqual([row["nombre"] for row in rows], ["Ana", "Beto"])

    def test_xlsx_to_csv_requires_headers(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no instalado")
        workbook = openpyxl.Workbook()
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "vacio.xlsx"
            target = Path(temp) / "vacio.csv"
            workbook.save(source)
            with self.assertRaises(pea.DataError):
                pea._xlsx_to_csv(source, target)

    def test_docx_extract_text(self):
        try:
            import mammoth
        except ImportError:
            self.skipTest("mammoth no instalado")
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "documento.docx"
            path.write_bytes(_minimal_docx("Hola mundo"))
            html = pea.extract_docx_html(path)
            self.assertIn("Hola mundo", html)


if __name__ == "__main__":
    unittest.main()
